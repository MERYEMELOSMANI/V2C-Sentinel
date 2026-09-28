from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'paper/working'
REF = WORK / 'IEEE_Access_Template_2024.docx'
OUT = ROOT / 'paper/V2C_Sentinel_IEEE_17_Page_Journal_Draft.docx'
d = Document(REF)
sections = [deepcopy(s._sectPr) for s in d.sections]
body = d._element.body
for child in list(body): body.remove(child)
source = []

STYLE_MAP = {
    'Body Text': 'PARA',
    'Heading 1': 'H1_List (Space)',
    'Heading 2': 'H2_Cont',
    'Heading 5': 'H1',
    'paper title': 'Paper Title',
    'Author': 'AU',
    'Keywords': 'IT',
    'figure caption': 'Fig Caption',
    'table title': 'Table Title',
    'table head': 'Table Title',
    'table copy': 'body type',
    'references': 'References',
}

def p(text, style='PARA'):
    para = d.add_paragraph(text, STYLE_MAP.get(style, style))
    source.append(text)
    return para

_first_subheading = True

def heading(text):
    global _first_subheading
    _first_subheading = True
    return p(text.upper(), 'H1_List (Space)')

def boundary(sect):
    para = d.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = Pt(1)
    para._p.get_or_add_pPr().append(sect)


def subheading(text):
    global _first_subheading
    style = 'H2_First' if _first_subheading else 'H2_Cont'
    _first_subheading = False
    return p(text.upper(), style)

def biography(name, first_paragraph, following_paragraph):
    para = d.add_paragraph(style='AU_Bios')
    para.paragraph_format.keep_together = True
    para.paragraph_format.keep_with_next = True
    name_run = para.add_run(name.upper())
    name_run.style = d.styles['AU_Bios bd']
    para.add_run(' ' + first_paragraph)
    source.append(name.upper() + ' ' + first_paragraph)
    follow = d.add_paragraph(following_paragraph, style='AU_Bios_No Space')
    follow.paragraph_format.keep_together = True
    source.append(following_paragraph)
    return para, follow

def figure(name,caption):
    para=d.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(5)
    para.paragraph_format.space_after=Pt(2)
    para.paragraph_format.keep_with_next=True
    shape=para.add_run().add_picture(str(ROOT/'paper/figures'/name),width=Inches(3.45))
    shape._inline.docPr.set('descr',caption)
    cap=p(caption,'figure caption')
    cap.paragraph_format.keep_together=True
    cap.paragraph_format.space_after=Pt(7)
    num=OxmlElement('w:numPr'); ident=OxmlElement('w:numId'); ident.set(qn('w:val'),'0'); num.append(ident)
    cap._p.get_or_add_pPr().append(num)

def equation(tokens,number):
    para=d.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(4)
    para.paragraph_format.space_after=Pt(6)
    math=OxmlElement('m:oMath')
    def run(text):
        r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=text; r.append(t); return r
    for token in tokens:
        if isinstance(token,tuple):
            sub=OxmlElement('m:sSub'); e=OxmlElement('m:e'); e.append(run(token[0])); v=OxmlElement('m:sub'); v.append(run(token[1])); sub.append(e); sub.append(v); math.append(sub)
        else: math.append(run(token))
    para._p.append(math); para.add_run('    ('+str(number)+')')
    source.append('Equation '+str(number)+': '+str(tokens))

def compact_table(caption, headers, rows, widths):
    cap = p(caption, 'table head')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    tbl = d.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    for cell, text_value in zip(tbl.rows[0].cells, headers):
        cell.text = str(text_value)
    for row_values in rows:
        for cell, text_value in zip(tbl.add_row().cells, row_values):
            cell.text = str(text_value)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge = OxmlElement('w:'+side)
        edge.set(qn('w:val'),'single')
        edge.set(qn('w:sz'),'4')
        edge.set(qn('w:color'),'808080')
        borders.append(edge)
    tbl._tbl.tblPr.append(borders)
    for i, row in enumerate(tbl.rows):
        trp = row._tr.get_or_add_trPr()
        trp.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trp.append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            for para in cell.paragraphs:
                para.style = d.styles['PARA']
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.space_after = Pt(1)
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j in (0, 2, 3) else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(7)
                    run.bold = i == 0
    source.append(caption+'\n'+' | '.join(headers)+'\n'+'\n'.join(' | '.join(map(str,r)) for r in rows))

title = 'V2C Sentinel Two Level CAN Masquerade Detection With Budget Limited Cloud Confirmation'
p(title, 'paper title')
p('Meryem EL OSMANI¹, Nissrine BESTOUT¹, Asmaa Berdigh²*, and Khalid EL YASSINI³*', 'Author').paragraph_format.space_after = Pt(4)
p('¹Master IASDO, Faculty of Sciences Meknes, Moulay Ismail University, Morocco', 'PI').paragraph_format.space_after = Pt(0)
p('²IA Laboratory, International University of Rabat, Morocco', 'PI').paragraph_format.space_after = Pt(0)
p('³Computer Science Department, Faculty of Sciences Meknès, Moulay Ismail University, Morocco', 'PI').paragraph_format.space_after = Pt(0)
p('*Corresponding authors: Asmaa Berdigh and Khalid EL YASSINI (email addresses to be confirmed).', 'PI').paragraph_format.space_after = Pt(8)

p('ABSTRACT—Cloud assistance can provide a second intrusion-detection decision, but communication delay and request limits affect when that decision becomes available. This paper examines a two-level alert architecture for controller area network masquerade detection. A local logistic regression model issues an immediate warning and, subject to a token-bucket budget, requests confirmation from a random forest model. A chronological replay combines automotive attack recordings with a separately recorded fifth-generation cellular delay trace. The evaluation contains one attack episode and one normal recording in each of the development and test splits. Local warnings identify both episodes after the assumed 5 ms processing time, while producing approximately 3568 and 3620 grouped false alerts per hour on the respective normal recordings. Cloud confirmation produces no observed false alerts on those recordings, but its timeliness depends on the request budget and recording. At 50 requests per second, confirmation takes 223.6 ms on development and 31.0 ms on test; unrestricted requests reduce these delays to 29.0 and 31.0 ms. The results support separating early warnings from confirmed alerts while showing that confirmation cannot be assumed to meet a 150 ms deadline. The study is exploratory and does not establish deployment performance or population-level detection rates.', 'Abstract')
p('INDEX TERMS—automotive security, cloud assistance, controller area network, detection latency, intrusion detection, masquerade attack, request budgeting', 'Keywords')
title_sect = deepcopy(sections[0])
boundary(title_sect)

heading('Introduction')
p('Controller area network (CAN) intrusion detection must distinguish malicious activity from legitimate vehicle traffic. The ROAD dataset includes masquerade attack traces designed to support this investigation [1]. Signal-level methods such as CANShield use relationships among signals to detect attacks that may be difficult to identify from message timing alone [2]. These studies motivate examining both the detection decision and the time at which an actionable warning becomes available.')
p('Consulting a second detector introduces a separate systems question. A local model can produce an early indication, whereas remote analysis requires a request, processing, and a returning response. Cellular measurements in the CICV5G dataset provide recorded vehicle–network–vehicle delays for studying this communication component [3]. A constrained request budget can introduce additional waiting before an informative message is selected. Consequently, classifier output alone does not determine episode-level alert latency.')
p('This paper asks how separating immediate local warnings from cloud-confirmed alerts affects false-alert burden and episode-level timeliness under limited request rates. V2C-Sentinel preserves each local positive decision as a level-1 warning and sends eligible messages for a level-2 decision. Confirmation is an additional output; a negative or delayed remote decision does not retract the earlier warning.')
p('The contribution is a reproducible replay case study of this separation, with a shared token-bucket mechanism and distinct measurements for the two alert levels. We report development and test recordings separately and expose the conditions under which confirmation misses an experimental deadline. We do not propose a new classifier or claim that remote execution is necessary. Both detectors execute on the development computer, and the communication channel is simulated from recorded delays.')

p('The evaluation distinguishes local-warning timeliness, false-alert burden at each alert level, and the confirmation budget needed to meet an episode-onset deadline. This prevents a low confirmation false-alert count from being mistaken for a low total warning burden.')


heading('Background and Design Motivation')
subheading('CAN Masquerade Attacks')
p('The controller area network is a broadcast bus designed for reliable real-time communication among electronic control units. Frames generally carry an identifier and payload but do not identify the transmitting device cryptographically. Receivers interpret the identifier as message priority and meaning. This architecture enables efficient control, yet it also means that a compromised unit can transmit frames using an identifier associated with another function.')
p('A masquerade attacker attempts to imitate legitimate traffic while altering the values or sequence that downstream components observe. Unlike flooding or obviously invalid identifiers, a masquerade may preserve nominal message frequency and use legal identifiers. Detection can therefore require relationships among signals, payload dynamics, temporal context, or consistency with other vehicle states. The ROAD dataset was designed in part to support evaluation of these difficult attacks [1].')
p('The term masquerade covers different experimental constructions. Some datasets inject fabricated frames; others replace fields within recorded traffic or post-process recordings. The operational realism, attack onset, affected signals, and labels vary. A detector result is meaningful only with the dataset construction and evaluation unit stated explicitly. This paper uses the ROAD signal-extraction files and does not claim to reproduce a live compromise.')
subheading('Local and Remote Detection Roles')
p('A local detector can observe the bus without depending on external connectivity. Its design is constrained by processor, memory, energy, integration, and certification requirements. A remote detector may use a larger model, aggregate broader context, or receive updates more easily, but it introduces communication and service dependencies. The appropriate split is therefore a systems choice rather than an automatic consequence of model accuracy.')
p('Cloud assistance can take several forms. Raw events may be offloaded for inference, local features may be transmitted, uncertain predictions may be deferred, or model updates may be exchanged through federated learning. These forms have different privacy, latency, bandwidth, and failure behavior. V2C-Sentinel studies per-event confirmation of locally positive messages. It does not offload training and does not aggregate a fleet model.')
p('Confirmation also differs from replacement. If the local output remains visible, cloud response changes confidence or escalation but not the earliest warning. If the local output is hidden, cloud response becomes the first actionable alert. A negative response may retract an earlier warning in some systems, whereas the present replay preserves it. These choices determine which latency and false-alert measures are relevant.')
subheading('Deadline Evaluation')
p('Detection latency can be measured from event observation, window completion, request transmission, or attack onset. Starting at request transmission excludes the time needed to select an informative event and can make a constrained policy appear faster than it is from the defender’s perspective. This paper starts at labeled episode onset because the research question concerns when the first episode-associated warning or confirmation becomes available.')
p('A deadline converts continuous delay into an application-facing outcome, but only after its origin and meaning are specified. The 150 and 300 ms values used here are experimental thresholds selected for comparison. They are not derived from a braking controller, powertrain requirement, or safety standard. A deployed system would need a hazard-specific deadline and an explicit response action.')
p('Deadline awareness can also refer to scheduling. A scheduler may predict whether a reply will arrive before a deadline and suppress requests expected to be late. The current system does not do that. It applies deadlines only during evaluation. Calling the admission policy deadline aware would therefore be inaccurate; the study is deadline evaluated and motivates deadline-aware routing as future work.')
subheading('False Alerts and Resource Budgets')
p('Frame-level false-positive rate does not directly express operator burden. CAN traffic can contain thousands of frames per second, so a small percentage may create a continuous warning stream. Grouping nearby positives produces an alert-level measure, but the answer depends on the grouping rule. The manuscript reports both raw positive counts and one-second grouped rates to keep that transformation visible.')
p('Communication budgets create another denominator. A request rate in requests per second can be compared across recordings of different duration, but it does not describe byte volume, energy, or monetary cost. A rate limiter also interacts with event timing: the same sustained capacity can admit different attack messages depending on preceding benign demand and initial bucket state.')
p('The central design tension is therefore three-way. A permissive local threshold can improve the chance of early episode-associated positives but create false warnings and heavy cloud demand. A strict threshold can reduce burden while preventing the cloud from seeing missed events. More cloud capacity can accelerate confirmation, but only if earlier admitted events are informative. The evaluation separates these effects instead of collapsing them into one accuracy value.')


heading('Related Work')
p('CAN Masquerade Detection and Timing', 'Heading 2')
p('ROAD [1] provides the data foundation for this study. Its masquerade traces are post-processed from recorded vehicle data, and should not be represented as organically executed masquerade attacks. CANShield [2] uses signal-level autoencoder ensembles and explicitly evaluates attack and event detection latency as well as hardware processing latency. Thus, measuring time from an attack to a detection is not itself a new contribution. Our focus is the additional timing and admission cost of obtaining a second model decision.', 'Body Text')
p('Moriano et al. [4] study signal-clustering similarity, Marfo et al. [5] use temporally enriched message-sequence graphs, and Moriano et al. [6] compare lightweight online methods. Together, they motivate separating classifier evidence, observation-window requirements, and response delivery rather than reporting one undifferentiated latency.', 'Body Text')
p('The MIDS preprint by Liu et al. [7] uses bidirectional Mamba processing for CAN identifiers and payloads and reports single-window inference latency. Guerra et al. [8] compare multiple intrusion detectors on ROAD and other automotive data. These results provide relevant detector and dataset context, but their performance figures cannot be transferred to the present classifiers, split, or hardware. We therefore do not place their reported accuracy beside our episode counts as if the experiments were directly comparable.', 'Body Text')
p('Rajapaksha et al. [9], Hoang and Kim [10], and Dong et al. [11] offer different local CAN detectors. Their results show that local model design and efficiency remain active alternatives to cloud assistance; remote inference should not be assumed necessary.', 'Body Text')
p('Vehicle Cloud Offloading and Selective Assistance', 'Heading 2')
p('Loukas et al. [12] investigate cloud-based cyber-physical intrusion detection using deep models on a small robotic vehicle. Their study includes a latency model for deciding when computation offloading is beneficial. This is a close architectural predecessor: cloud-based vehicle intrusion detection and latency-aware offloading are established ideas. V2C-Sentinel differs in evaluating selected CAN masquerade recordings, preserving an immediate local warning, and measuring the first confirmed alert under a finite request-admission budget.', 'Body Text')
p('Chinchali et al. [13] formulate cloud-robotics offloading as a sequential decision problem balancing performance against communication cost. The present fixed local-positive admission rule and token bucket do not learn when cloud assistance is beneficial; this is a constrained replay evaluation, not a general routing method.', 'Body Text')
p('Reliability Routing and Federated Detection', 'Heading 2')
p('Azam et al. [14] present a reliability-aware edge–cloud intrusion-detection framework that uses calibrated packet-prefix evidence to choose local decisions, continued observation, or cloud refinement. This is a close conceptual comparison because cloud access is selective. V2C-Sentinel does not implement calibrated abstention or a continue-observing action. It forwards locally positive CAN events when tokens are available and then evaluates confirmation against elapsed-time deadlines. The difference is in routing semantics and evaluation setting, not in the existence of an edge–cloud structure.', 'Body Text')
p('Mao et al. [15] formalize cost-aware routing among experts. Heidari et al. [16] select participants for collaborative vehicle-model training. Neither mechanism is equivalent to this fixed, per-message cloud-confirmation replay.', 'Body Text')
p('Scope of the Present Contribution', 'Heading 2')
p('CICV5G [3] supplies an independently recorded cellular delay trace. Combining that trace with ROAD allows a controlled study of how a request budget changes confirmation time while the local output remains available. The evidence supports a small, deterministic case study rather than a first-of-its-kind cloud detector. In particular, the deadline is used to evaluate alerts; the implemented request policy does not predict whether a reply will meet that deadline. Comparative accuracy, routing optimality, and deployment safety remain outside the demonstrated claims.', 'Body Text')

subheading('Evidence Synthesis and Research Gap')
p('The literature evidence workbook supplied for this study records detector location, cloud role, delay treatment, deadline use, false-alarm measurement, evaluation granularity, baselines, and limitations for the selected sources. That structure reveals that superficially similar papers often answer different questions. A method may be cloud based without deciding when to offload, delay aware without enforcing an alert deadline, or selective without routing individual detections. The present paper therefore compares mechanisms rather than grouping studies by title keywords alone.')
p('The reviewed CAN literature establishes several strong local alternatives. Signal reconstruction, clustering, graph learning, sequence modeling, adversarial autoencoding, and hidden Markov models all aim to expose attacks that preserve plausible identifiers or timing. Their evaluation units vary among frames, windows, events, and complete recordings. A reported inference time is consequently not equivalent to episode-onset detection delay: the detector may first require a window to fill, a signal to deviate, or a sequence statistic to stabilize.')
p('Cloud and edge studies introduce a second distinction. Loukas et al. [12] model vehicle offloading latency, whereas Chinchali et al. [13] learn network offloading decisions for robotic perception. Azam et al. [14] route uncertain packet-prefix decisions between continued observation, local output, and cloud refinement. These systems demonstrate that computation placement and communication cost are established concerns. The remaining question addressed here is narrower: what happens when a noisy local positive is preserved as a warning while cloud admission is rate limited and confirmation is judged against an episode deadline?')
p('Learning-to-defer methods [15] formalize expert selection and the cost of requesting a second decision. Their guarantees depend on training objectives and data assumptions not implemented here. Likewise, the federated edge-cloud architecture in [16] selects vehicles for training rounds; it does not send a suspicious CAN event to an independent cloud classifier. Treating either study as the same mechanism would overstate both the novelty and the maturity of V2C-Sentinel.')
p('Across the evidence table, false alarms and latency are frequently reported, but seldom with the same definitions used in this replay. Some papers count misclassified frames, others count attack windows, and others report end-to-end service latency. A false-positive percentage cannot be converted into alerts per hour without traffic exposure and a grouping rule. A network round-trip statistic cannot be converted into episode detection time without knowing when an informative event was selected. These incompatibilities prevent a responsible cross-paper numerical ranking.')
p('The literature instead supports three design requirements. First, a local baseline must remain visible because cloud confirmation cannot recover events that are never requested. Second, timing must start at attack onset rather than at request transmission if the intended claim concerns response. Third, communication demand must be reported with the alert result because a low false-confirmation count achieved by querying nearly every positive event may still be operationally expensive.')
p('This synthesis also narrows the contribution. V2C-Sentinel is not presented as the first vehicle-cloud detector, the first timing-aware intrusion detector, or the first selective offloading system. Its contribution is an auditable separation of two alert semantics under replayed cellular delay and explicit request admission, together with evidence showing why a favorable confirmation result cannot erase a noisy local-warning channel.')
p('The selected literature remains a focused set rather than a systematic review. It was assembled to cover ROAD-based masquerade detection, local detector alternatives, vehicle-cloud offloading, reliability routing, learning to defer, and federated edge-cloud detection. It does not support bibliometric claims about the prevalence of any method. The comparison is used to establish mechanism-level distinctions and to prevent unsupported novelty statements.')

compact_table('TABLE I. REPRESENTATIVE MECHANISM COMPARISON',
    ['Study', 'Local CAN detector', 'Cloud role', 'Timing evidence'], [
    ['CANShield [2]', 'Signal autoencoders', 'None', 'Event and processing latency'],
    ['Moriano et al. [4]', 'Signal clustering', 'None', 'Observation dependent'],
    ['Marfo et al. [5]', 'Graph learning', 'None', 'Temporal graph context'],
    ['Dong et al. [11]', 'Multi-observation HMM', 'None', 'Window sensitivity'],
    ['Loukas et al. [12]', 'Vehicle IDS component', 'Inference offload', 'Latency model'],
    ['Chinchali et al. [13]', 'Task dependent', 'Learned offload', 'Communication cost'],
    ['Azam et al. [14]', 'Calibrated local evidence', 'Selective refinement', 'Continued observation and cloud'],
    ['Heidari et al. [16]', 'Vehicle GRU', 'Federated aggregation', 'End-to-end alert latency'],
    ['V2C-Sentinel', 'Logistic warning gate', 'Per-event confirmation', 'Onset deadline and request budget'],
    ], [0.78, 0.83, 0.82, 1.02])

subheading('Comparison by Detection Representation')
p('Signal-level methods seek relationships that should remain consistent during legitimate operation. CANShield [2] uses autoencoder ensembles over signal groups and evaluates both attack and event detection latency. Signal-clustering similarity [4] instead monitors how relationships among decoded signals change. These methods are especially relevant to masquerade attacks because identifier frequency alone may remain plausible. They also show that window formation and event definition contribute to latency before any network delay is added.')
p('Graph-based work [5] represents message sequences and temporal attributes as structured relationships. This can capture dependencies that flat per-frame classifiers miss, but it introduces graph construction and context requirements. The present logistic regression does not model those relationships. Its purpose is to provide a lightweight, transparent gate whose operational weaknesses can be measured, not to claim state-of-the-art masquerade accuracy.')
p('Sequence and representation-learning approaches [7], [9], [10] can learn temporal or latent structure from identifiers and payloads. Their reported inference costs and window definitions differ, and their accuracy is evaluated on different splits and datasets. They support the premise that a stronger local detector may be possible. For this reason, cloud assistance must eventually be compared with improved local models rather than justified by the weakness of one thresholded baseline.')
p('The multiple-observation hidden Markov model in [11] is notable because it evaluates masquerade attacks on a real-vehicle platform and classifies frames using surrounding-window statistics. It reports strong performance while identifying sophisticated attacks as a continuing challenge. This combination makes it a useful local-baseline candidate and illustrates why frame-by-frame output can still depend on multi-frame context.')
p('Guerra et al. [8] compare multiple intrusion detectors on ROAD and other automotive datasets. Comparative studies are valuable because model rankings can change with attack and data representation. Their results cannot be copied into the present paper as baselines because the feature pipeline, split, threshold, and evaluation units differ. A fair comparison would rerun candidate models under the frozen V2C-Sentinel protocol.')
subheading('Comparison by Cloud Function')
p('In the cloud-based vehicle intrusion work of Loukas et al. [12], offloading is motivated by computational capability and evaluated with a latency model. This establishes that vehicle intrusion detection and cloud computation have been combined before. The present distinction lies in preserving a local warning, applying an explicit request budget, and measuring confirmation from attack onset with replayed communication delay.')
p('Network offloading for cloud robotics [13] treats routing as a sequential decision that balances task performance and communication cost. That formulation is broader and more adaptive than the fixed token bucket used here. It provides a conceptual baseline for future learned admission, but applying it to intrusion confirmation would require a state representation, reward, safe exploration strategy, and enough independent episodes for training and validation.')
p('Reliability-aware edge-cloud routing [14] is a close comparison because it chooses among local decisions, continued observation, and remote refinement. Its calibrated evidence and packet-prefix setting differ from the present local-positive rule. The comparison suggests two useful extensions: an explicit abstention region and a continue-observing action that can avoid premature warnings without immediately consuming cloud capacity.')
p('The learning-to-defer framework in [15] formalizes routing among experts under cost. It emphasizes that the router and experts should be evaluated jointly. V2C-Sentinel does not train a router and cannot inherit those theoretical guarantees. Its budget sweep instead exposes the behavior of one hand-specified gate and supplies a baseline against which a learned deferral policy could be tested.')
p('The hierarchical federated vehicle IDS in [16] uses edge and cloud resources for collaborative training, participant selection, and explainability. Its DRL selector chooses which vehicles contribute to training rounds, not which suspicious event receives a second prediction. It is relevant to architecture and communication realism, but it does not close the per-event confirmation question studied here.')
subheading('Comparison by Timing Evidence')
p('Timing evidence in the selected papers falls into at least four categories: model inference time, window or observation time, communication or service latency, and end-to-end alert latency. These categories answer different questions. A submillisecond inference does not guarantee prompt detection if a long context window is required, and a short network delay does not guarantee prompt confirmation if an informative request is admitted late.')
p('CANShield [2] explicitly reports event and processing latency, while the online lightweight evaluation in [6] exposes detection-computation trade-offs. The cloud and edge studies [12], [14], [16] model or measure communication-related delay. Together they motivate an additive timing decomposition, but the components cannot be combined numerically across papers because their start events, hardware, traffic, and endpoints differ.')
p('The current equations separate local completion, request delay, and cloud processing, yet the observed episode confirmation includes another implicit term: time from attack onset until the first admitted event that yields a positive forest result. That selection term explains why confirmation can exceed the network delay by a large margin under tight budgets. It should be made explicit in future per-request logs.')
p('No reviewed paper provides a directly transferable safety deadline for this experiment. Some report latency that they describe as real time, but real-time adequacy depends on the controlled function and response. The manuscript therefore treats 150 and 300 ms as sensitivity points and avoids using them as evidence of compliance with an automotive safety requirement.')
p('The research gap is consequently empirical and definitional. The paper evaluates one transparent warning-confirmation cascade with explicit alert semantics, onset-based timing, request admission, and separate burden measures. Its value lies in showing how those definitions change the interpretation of apparently successful detection, while its limited sample prevents claims about general superiority.')

heading('System and Replay Method')

subheading('Threat Model and Alert Semantics')
p('The replay assumes an adversary capable of generating masquerading CAN traffic represented by the selected ROAD labels. The attacker can alter the content associated with legitimate traffic; the study does not simulate initial compromise, cryptographic authentication, or a live attacker. The detector observes recorded identifiers and extracted signals. Model files, local processing, and the conceptual cloud service are trusted. Attacks on the confirmation channel, adversarial model manipulation, and intentional denial of the request budget are outside the evaluated threat model.')
p('A level-1 warning means that the local score crossed its operating threshold. A level-2 confirmation means that an admitted message received a positive forest prediction. The latter is a model output, not proof of an attack. Absence of confirmation can mean that the message was not admitted, that its reply has not arrived, or that the forest predicted benign traffic. Fig. 1 shows these separate paths and the persistence of the local warning.')
figure('fig1_architecture.png','Fig. 1. Two-level replay architecture. Positive local predictions produce warnings immediately and compete for cloud admission. Recorded communication delays affect only the confirmation path.')
subheading('Feature Processing and Fixed Classifiers')

p('Each input event contains a timestamp, message identifier, extracted signal values, and an evaluation label. The label is used only for training or evaluation as appropriate to the split. The prediction pipeline uses the identifier and Signal_ columns as features, replaces missing values with −1, and standardizes features using statistics fitted on the training split. Timestamp and label columns are excluded from the feature matrix.')
p('The local classifier is logistic regression with a maximum of 1000 iterations and random seed 42. The remote classifier is a random forest with 50 trees, maximum depth 15, and random seed 42. Both classifiers are trained on the normal and attack training recordings. The local positive threshold is the 95th percentile of attack-class probabilities on the normal development recording; equality to the threshold is classified as positive. The forest uses its predicted class. This local threshold is distinct from the suspicion threshold used in an earlier policy-selection experiment.')
subheading('Admission Control and Response Timing')
p('For a message observed at time t, local processing completes at t + 5 ms. A positive local prediction immediately produces a level-1 warning. Only locally positive messages are eligible for a cloud request. A token bucket replenishes at b requests/s, has capacity max(2, 0.1b) requests, and starts full for each recording. One admitted request consumes one token. Budgets of 1, 5, 10, and 50 requests/s are compared with an unrestricted condition. Unrestricted means that every locally positive message is requested, rather than every CAN message.')
p('An admitted request returns after the local completion time plus the replayed communication delay and 10 ms of assumed cloud processing. Positive forest predictions produce level-2 alerts at their return times. Replies are ordered chronologically, and replies outstanding at the end of a recording are processed. A negative forest prediction produces no level-2 alert. The simulator does not model a service queue, packet loss, retransmissions, or concurrent inference capacity.')

p('Let tᵢ denote the observation time of message i, τL the assumed local processing time, τC the assumed cloud processing time, and d(s) the replay delay at request time s. For a locally positive message, the warning time is given by (1). If that message is admitted, its confirmation time, conditional on a positive forest prediction, is given by (2).')
equation([('s','i'),' = ',('t','i'),' + ',('τ','L')],1)
equation([('c','i'),' = ',('s','i'),' + d(',('s','i'),') + ',('τ','C')],2)
p('Equations (1) and (2) show why an episode-level deadline is stricter than a per-request response target: a suitable message may be admitted only after the attack has already been active. There is no waiting queue for rejected requests. An eligible message without a token is skipped, and a later local positive may obtain the next token. Thus the observed confirmation delay combines event selection after onset with the delay of the selected reply.')

subheading('Delay Replay and Deadline Interpretation')
p('Trace timestamps are measured relative to the first trace sample. At each request time, the simulator wraps time by the trace duration and selects the first trace sample at or after the wrapped time. Replayed delays have a 1 ms minimum. The recorded delay is used directly as the communication-delay term; the two datasets have no common time alignment. Whether the measurement includes application processing that overlaps the added 10 ms term requires calibration before an end-to-end deployment interpretation.')
p('The reporting deadlines are 150 and 300 ms from attack onset. They classify observed detection times and do not suppress replies or change scheduling. In particular, the local warning is not held until a deadline, and late confirmations remain available as diagnostic information. These deadlines are experimental settings, not validated automotive safety requirements.')

subheading('System Invariants and Failure Modes')
p('The simulator preserves four invariants. A local warning is emitted before any request decision; cloud eligibility requires a local positive; one accepted request consumes one token; and a negative or late cloud result never retracts the earlier local warning. These rules make the two output channels auditable. They also prevent the evaluation from silently crediting cloud assistance for a warning that was already available locally.')
p('The first invariant separates latency from confidence. The 5 ms local time is an assumed processing constant, so every locally detected episode has the same first-warning delay. Cloud confirmation can add evidence, but it cannot reduce that already-recorded delay. A paper that reports only the confirmed alert would describe a different interface in which local warnings are hidden or nonactionable. The present results must be interpreted under the explicit two-level interface.')
p('The eligibility rule creates a cascade dependency. If logistic regression produces a false negative for every attack event, the random forest is never consulted for that episode. If it produces many false positives, benign events compete for the same tokens as attack events. The local detector therefore influences both safety coverage and communication demand, even though cloud predictions determine level-2 alerts.')
p('The token bucket represents sustained request capacity rather than a per-second counter. At rate b, fractional tokens accumulate continuously up to the configured capacity. A recording starts with a full bucket, allowing a small initial burst. This model captures a common rate-control mechanism but omits byte size, transport congestion, retransmission, and server service capacity. Two request streams with the same count could impose different network load if their payloads differ.')
p('Rejected events are skipped rather than queued. This choice bounds backlog but makes confirmation dependent on the timing of later local positives. A queue would produce different behavior: it could preserve potentially informative requests while adding waiting time and possibly processing stale evidence. A coalescing policy could instead replace a pending request with newer evidence. Those alternatives are outside this experiment and should not be inferred from the token-bucket results.')
p('The replay also assumes that every accepted request receives the recorded communication delay independently of other requests. It does not enforce a maximum number of simultaneous cloud inferences. Under unrestricted admission, many requests can therefore be in flight without increasing service time. This optimistic concurrency assumption matters most where the local detector is noisy and the request rate exceeds 100 requests/s.')
p('A level-2 false-alert count of zero can result from a specific combination of admission and classifier output. At low budgets, fewer benign local positives reach the forest; at high budgets, the forest still predicts those accepted benign events as normal in the observed recordings. The experiment does not distinguish whether the same behavior would persist under a different vehicle, temperature, attack family, or network phase.')
p('Late confirmations remain diagnostically meaningful in the logs. They may help forensic review or support a slower response process, but they do not satisfy the tested deadline. Counting eventual detection and timely detection separately prevents late results from being credited as prompt response. This distinction is especially important at the 1 request/s budget, where both evaluated episodes are eventually confirmed but only after delays far beyond 300 ms.')
p('The architecture produces no physical control command. Warning semantics, operator escalation, automated mitigation, and fail-safe behavior are application decisions that require their own hazard analysis. The experiment therefore evaluates alert generation and communication timing only. Terms such as safe, real time, and deployable are avoided unless tied to an explicit measured condition.')

heading('Experimental Setup and Metrics')
subheading('Data Partition and Experimental Conditions')
p('The ROAD training pair consists of the basic-long ambient recording and correlated-signal masquerade recording 1. Development uses the radio-infotainment ambient recording and masquerade recording 2. Test uses the winter ambient recording and masquerade recording 3. These correspond to normal_01 and attack_01, normal_02 and attack_02, and normal_03 and attack_03 in the saved split manifest. The cellular replay uses the CICV5G urban-road n78 trace. Table II reports the evaluated recordings; the durations are the last timestamp minus the first timestamp.')

p('The split is chronological at the recording level rather than a random frame split. Frames from one recording therefore remain together, reducing direct leakage through near-duplicate temporal neighbors. The training split supplies both normal and attack examples for the two supervised classifiers. The normal development recording supplies the local operating threshold, and the development attack recording is used for exploratory interpretation. The test recordings were not used for fitting, but prior inspection means they cannot support a claim of pristine holdout evaluation.')
p('ROAD signal-extraction files expose a message identifier, timestamp, label, and a variable set of decoded signal columns. The pipeline constructs the feature vector from the identifier and columns whose names begin with Signal_. Missing signal positions are filled with -1 before scaling. This sentinel distinguishes missing fields from decoded zeros, although it may also become an artificial cue associated with message type. A deployment pipeline should verify that signal availability and decoding conventions match training.')
p('The standardizer is fitted on the combined normal and attack training frames and reused for development and test. Timestamp and label are explicitly excluded from the feature matrix. A guard raises an error if either appears among the selected features. This check was added after an earlier leakage issue and is part of the current scientific provenance. The reported results should not be combined with preliminary tables produced before the corrected pipeline was established.')
p('Both models are supervised in the present pipeline. Logistic regression is used locally because it is compact and inexpensive on the development computer. The random forest provides a nonlinear second decision. The names describe computational roles in the replay, not validated vehicle and cloud deployments: both model files are executed on the same computer, and their processing times are assigned by the simulator.')
p('The local threshold is the 95th percentile of attack-class probabilities on the normal development recording. This intentionally yields an event-level positive fraction near 5% on that recording before temporal grouping. It is an operating-point choice rather than a learned safety requirement. The resulting burden illustrates why threshold selection must be evaluated in alerts per unit time and not only in frame-level percentages.')
p('The forest uses its predicted class without probability recalibration. The experiment therefore compares a thresholded probabilistic local detector with a class-output cloud detector. It does not match the two models at equal false-positive rate, equal recall, equal compute, or equal calibration. Any claim that the forest is intrinsically stronger would require those controlled comparisons and independent data.')
p('The two attack recordings contribute one continuous episode each under the 0.5 s gap rule. Thousands of positive frames within an episode are correlated observations and do not create thousands of independent attacks. This distinction governs reporting: the manuscript gives episode counts of 0/1 or 1/1 and reports frame counts only to explain request and false-alert behavior.')
p('Normal exposure is also limited. Development supplies 390.456 s and test supplies 47.731 s of normal traffic. Their combined 438.187 s is sufficient to expose a severe local warning burden, but too short to establish a rare-event cloud false-alert rate. Longer normal drives across operating conditions are required before estimating workload or a confidence interval for rare confirmed alerts.')
p('The 5G trace is an external measurement source rather than a synchronized companion recording. Simulation time is wrapped over the trace duration, and the first sample at or after the wrapped request time supplies the delay. This construction preserves measured delay values but not causal alignment between a vehicle event and network conditions. Trace-phase sensitivity is therefore a required extension.')

def table(caption, headers, rows, widths):
    cap = p(caption, 'table head')
    # Explicit caption text avoids inheriting sample automatic numbering.
    n = OxmlElement('w:numPr'); v=OxmlElement('w:numId'); v.set(qn('w:val'),'0'); n.append(v); cap._p.get_or_add_pPr().append(n)
    cap.paragraph_format.keep_with_next = True
    tbl = d.add_table(rows=1, cols=len(headers)); tbl.autofit=False
    tbl.alignment=1
    for col,w in zip(tbl.columns,widths): col.width=Inches(w)
    for cell,text in zip(tbl.rows[0].cells,headers): cell.text=text
    for row in rows:
        for cell,text in zip(tbl.add_row().cells,row): cell.text=str(text)
    pr=tbl._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'000000'); borders.append(edge)
    pr.append(borders)
    for i,row in enumerate(tbl.rows):
        trp=row._tr.get_or_add_trPr(); trp.append(OxmlElement('w:cantSplit'))
        if i==0: trp.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j])
            for para in cell.paragraphs:
                para.style=d.styles['PARA']; para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before=Pt(2); para.paragraph_format.space_after=Pt(2)
                for run in para.runs:
                    run.bold=(i==0)
                    run.font.size=Pt(8)
    source.append(caption+'\n'+' | '.join(headers)+'\n'+'\n'.join(' | '.join(map(str,r)) for r in rows))

table('TABLE II.  EVALUATION RECORDINGS', ['Recording','Messages','Duration (s)','Episodes'], [
    ['Development attack', '63 260','28.227','1'], ['Development normal','874 018','390.456','0'],
    ['Test attack','38 003','16.964','1'],['Test normal','106 939','47.731','0']], [1.17,.75,.76,.65])
subheading('Episode and Alert Metrics')
p('An attack episode groups positive-labeled events whose successive timestamps are separated by no more than 0.5 s. Each evaluated attack recording contains one such episode. Detection delay is the earliest warning time linked to a true-positive event in that episode minus episode onset. A detection is timely when this delay is no greater than the reporting deadline. Because each split contains one episode, timely results are reported as detected episode counts rather than as population estimates.')

p('For episode e with onset aₑ and alert level k, let Wₑ,ₖ be the set of alert times linked to positive-labeled events in that episode. Its first-detection delay is defined by (3), with the minimum of an empty set treated as infinity. Timeliness at deadline D is the condition Δₑ,ₖ ≤ D. This definition credits an alert associated with the episode even if the reply arrives after the episode has ended; its delay is still measured from onset.')
equation([('Δ','e,k'),' = min{w − ',('a','e'),': w ∈ ',('W','e,k'),'}'],3)

p('False-alert rates use separate normal recordings. The first false warning starts a group; a new group begins when a later warning is more than 1 s after that group start. Grouped count divided by recording duration gives grouped false alerts per hour. Dense warnings can approach one group per second, and short-recording boundaries can yield a rate slightly above 3600 per hour.')
p('Request totals are summed over the attack and normal recordings within each split, once per budget. They are not summed across deadlines because the same events are evaluated at both deadlines. Raw totals should not be compared across splits without accounting for their different durations. The token bucket permits an initial burst, so a nominal budget specifies a sustained admission rate rather than a strict cap in every one-second interval.')

p('For exposure normalization, Nq denotes admitted requests and T denotes the combined duration in seconds of the normal and attack recording within a split. The pooled request rate is Rq = Nq/T. This is a ratio of totals, not an unweighted average of per-recording rates. Development has 418.683 s of combined exposure and test has 64.694 s. Reporting pooled rates makes communication demand comparable despite the large difference in raw recording length.')

subheading('Metric Interpretation')
p('Timely detection is evaluated per episode and alert level. If no associated alert exists, the delay is infinite and the episode is missed. If an alert exists after the deadline, the episode is eventually detected but not timely. This three-state interpretation separates classification coverage from response timing and prevents a late positive from being counted as an on-time outcome.')
p('Association uses the true label of the originating event. A level-2 reply may arrive after the labeled attack interval ends and still be linked to that episode because the request originated from an attack event. This choice measures when evidence about the episode becomes available. A stricter operational metric could reject replies after the episode or after a mitigation opportunity closes, but that requires application semantics not available here.')
p('False alerts are measured on recordings with no positive labels. This avoids ambiguity about benign predictions occurring inside an attack recording, where background traffic and attack frames are interleaved. It also means the reported normal burden does not capture collateral false positives during an attack episode. A fuller evaluation should report both clean-session and within-attack false-alert behavior.')
p('The one-second grouping window approximates repeated-warning suppression. Because the window is anchored at its first warning, a dense stream can create a new group shortly after each second. A sliding extension rule would merge a continuous stream into fewer groups. Neither definition is intrinsically correct; the important requirement is that the rule be fixed and reported with the rate.')
p('Request utilization can be expressed relative to the nominal budget for finite conditions, but an initial burst and finite duration can produce small deviations. The manuscript therefore reports actual requests per second rather than interpreting every nominal setting as an exact realized rate. Unrestricted admission has no nominal denominator and is reported directly.')
p('No aggregate utility score combines timeliness, false alerts, and requests. Such a score would require application-specific weights and could conceal unacceptable behavior in one dimension. The paper presents the dimensions separately so that a reader can apply a relevant operating constraint. Future work may use Pareto analysis once enough independent cases exist.')
p('The absence of confidence intervals is deliberate. The deterministic budget rows reuse the same recordings, predictions, and trace, while deadline rows relabel the same alert times. Treating them as repeated samples would understate uncertainty. The most meaningful uncertainty is between recordings, vehicles, attacks, and network conditions, none of which is adequately sampled here.')
p('All numeric precision is limited to what the simulation supports. Millisecond values are shown to one decimal where they distinguish conditions, but this formatting does not imply measurement accuracy at that resolution. Processing constants are assumptions and trace alignment is synthetic. Conclusions rely on large qualitative differences and explicit threshold crossings rather than fine-grained timing precision.')
subheading('Reproducibility and Evidence Boundaries')
p('The paper uses only `results/final_two_level/`, whose 40 recording-level evaluations cover two splits, two recordings per split, five budgets, and two reporting deadlines. A deadline change reclassifies the same alert trajectory and is not an independent experiment. The deterministic replay therefore reports observed values without confidence intervals.')

p('The test recordings were inspected during earlier development. The recorded protocol states that they were not used for model fitting or threshold selection, but they are not an untouched holdout. The present two-level experiment is a later analysis than the earlier frozen multi-policy protocol. Accordingly, these results should be treated as exploratory. Processing times of 5 and 10 ms are simulation assumptions, not measurements on vehicle hardware.')

subheading('Experimental Matrix and Controls')
p('Each split contains one normal and one attack recording. For each budget, the simulator processes both recordings with the same classifier outputs, threshold, token-bucket rule, network trace, and processing constants. The 150 and 300 ms deadlines are applied after the trajectory is generated. They therefore create two labels for one alert history rather than two independent simulations. Request totals are reported once per budget to avoid double counting.')
p('The unrestricted condition is a diagnostic upper bound on admission among locally positive events. It does not query every CAN frame. Comparing it with finite budgets isolates the effect of admission under the same local eligibility rule. It does not isolate the value of the local gate itself, because an always-query baseline and a cloud model executed on every frame are not included in the two-level result set.')
p('The local path is invariant to budget by construction. This property serves as a simulator check: any budget-dependent local warning time would indicate unintended coupling. Likewise, level-2 false alerts can occur only after a local positive, successful admission, reply arrival, and positive forest result. The result table was checked against these causal constraints before manuscript generation.')
p('Development and test are reported separately because their exposures, traffic composition, candidate rates, and observed confirmation delays differ substantially. Pooling the two episodes into a 2/2 rate would hide that the same budget can miss a development deadline and meet a test deadline. Separate reporting makes this instability visible, although it cannot estimate its population frequency.')
p('The experiment does not repeat the fixed trace at multiple offsets or random seeds. It is deterministic for the saved predictions and initial trace phase. Error bars would therefore be misleading: repeated execution would reproduce the same points. Uncertainty must be obtained through new recordings, trace phases, bootstrapping units that respect temporal dependence, or a prespecified stochastic model rather than by duplicating the current run.')
p('The saved summary and all-runs files under results/final_two_level are the numerical source of truth for the manuscript. The older results/final directory previously contained outputs from incompatible scripts and has been quarantined. This separation prevents a configuration from one experiment being presented beside a summary from another. The paper does not use the legacy mixed output set.')
p('Reproduction requires the split manifest, trained model files, standardizer, per-frame predictions, delay trace, experiment script, and configuration constants. A complete archival release should additionally record hashes, Python and library versions, operating system, command line, and generated file checksums. Those items are identified as release requirements because the current repository does not yet provide a fully pinned execution environment.')
p('The saved development benchmark is contextual rather than part of the replay grid. It measures inference on the available computer and reports approximately 0.22 ms per local prediction and 16.0 ms per forest prediction for single messages. The replay nevertheless assumes 5 and 10 ms. The mismatch is disclosed rather than adjusted retrospectively because changing the constant would define a new experiment.')

heading('Results')
subheading('Early Warnings and Confirmation Timeliness')
p('Level 1 produces its first episode-associated warning after 5.0 ms in each split, satisfying both reporting deadlines. This does not wait for cloud admission or response. The normal development and test recordings produce 3568.1 and 3620.3 grouped local false alerts per hour, respectively, showing that the first level remains persistently noisy.')
p('Fig. 2 compares first-detection delays, and Table III reports the corresponding deadline outcomes. At the 150 ms deadline, none of the four finite budgets confirms the development episode on time. The 50 requests/s condition becomes timely when the deadline increases to 300 ms. On test, 5, 10, and 50 requests/s meet both deadlines, while 1 request/s misses both. Unrestricted requests meet both deadlines in both splits. Every listed condition eventually confirms its evaluated episode. In budget order 1, 5, 10, 50, and unrestricted, development delays are 3207.9, 805.7, 703.6, 223.6, and 29.0 ms; test delays are 725.8, 129.8, 129.8, 31.0, and 31.0 ms.')
figure('fig2_latency.png','Fig. 2. First-detection delay versus request budget. Each confirmation series contains one attack episode per split. The dotted line is the assumed 5 ms local path; horizontal lines mark 150 and 300 ms. The vertical axis is logarithmic.')
table('TABLE III.  TIMELY CLOUD DETECTIONS OUT OF ONE EPISODE', ['Budget\n(req/s)','Dev\n150 ms','Dev\n300 ms','Test\n150 ms','Test\n300 ms'],[
    ['1','0/1','0/1','0/1','0/1'],['5','0/1','0/1','1/1','1/1'],['10','0/1','0/1','1/1','1/1'],['50','0/1','1/1','1/1','1/1'],['Unrestricted','1/1','1/1','1/1','1/1']], [1.05,.57,.57,.57,.57])
p('The absence of a test-delay improvement between 5 and 10 requests/s shows that additional requests do not necessarily select an earlier informative event in a particular replay. Conversely, increasing the development budget to 50 requests/s improves confirmation delay substantially but still does not meet 150 ms. These are recording-specific observations; the experiment does not isolate the causal contributions of token availability, prediction sequence, and trace phase.')
subheading('False Alerts and the Meaning of Confirmation')
p('No level-2 false alerts are observed on either normal recording at any budget. The total normal exposure is approximately 438.2 s, or 7.3 min. Zero observed alerts over this short exposure does not establish a zero underlying false-alert rate. In addition, level-1 warnings remain present, so the combined architecture does not eliminate the local warning burden. An operational reduction in operator workload would depend on how the two outputs are presented and acted upon.')

figure('fig3_false_alerts.png','Fig. 3. False-alert burden on the separate normal recordings. Rates are identical across tested budgets. Zero confirmed alerts were observed over only 438.19 s combined exposure; local warnings remain present.')
p('The normal development file contains 44 288 positive local predictions among 874 018 events, while the normal test file contains 5068 among 106 939 events. These correspond to event-level positive fractions of approximately 5.07% and 4.74%. The grouping rule transforms these dense warning streams into 387 and 48 grouped alerts, respectively. These counts explain why the grouped rates remain near one alert per second. Event-level false positives and grouped operator alerts are therefore different quantities and should not be used interchangeably.')
subheading('Communication Demand')

table('TABLE IV.  CLOUD REQUEST TOTALS BY SPLIT', ['Budget\n(requests/s)','Development','Test'],[
    ['1','422','67'],['5','2095','326'],['10','4186','650'],['50','20889','3071'],['Unrestricted','49539','6503']], [1.13,1.1,1.1])
p('Table IV combines requests from the attack and normal recording in each split. These counts measure communication demand in requests, not bytes, monetary cost, or link utilization. The normal traffic consumes tokens as well as producing false local warnings. Consequently, sustained false positives can compete with attack-related messages for confirmation capacity.')

figure('fig4_request_rates.png','Fig. 4. Actual request rates pooled over the normal and attack recording within each split. Initial token bursts and finite exposure allow small deviations from the nominal sustained budget. Unrestricted admission still queries only local positives.')
p('In the unrestricted condition, development sends 49 539 requests over 418.683 s, approximately 118.3 requests/s, and test sends 6503 over 64.694 s, approximately 100.5 requests/s. At a 50 requests/s budget, the totals correspond to approximately 49.9 and 47.5 requests/s. The smaller test rate is consistent with candidate availability as well as admission limits; the budget is not a requirement to transmit at that rate. A deployed design would also need a byte budget and a server-capacity limit.')
p('The requests demonstrate a trade-off rather than a uniformly favorable budget. Development requires unrestricted admission to meet 150 ms, whereas test meets that deadline with 5 requests/s. Selecting a recommended budget from the favorable test recording would overstate the evidence. A useful configuration must instead be selected against independently measured operating requirements and a broader distribution of attack onsets and benign demand.')

subheading('Budget Response and Recording Dependence')
p('The delay curves are monotone nonincreasing within each split, but they are not smooth. Development improves from 3207.9 ms at 1 request/s to 805.7 ms at 5 requests/s, then only to 703.6 ms at 10 requests/s. The next increase to 50 requests/s produces a much larger change to 223.6 ms. This stepwise behavior is expected when the metric is the first confirmed event: additional capacity matters only when it admits an earlier event that the forest classifies as positive.')
p('Test shows a different step pattern. Increasing the budget from 1 to 5 requests/s changes the first confirmation from 725.8 to 129.8 ms, whereas 10 requests/s produces the same 129.8 ms result. At 50 requests/s the delay falls to 31.0 ms, matching the unrestricted condition. The plateau between 5 and 10 requests/s is evidence that nominal capacity alone does not determine latency in one recording.')
p('The unrestricted development delay of 29.0 ms is 2 ms shorter than the unrestricted test delay of 31.0 ms. These values include the assumed 5 ms local processing and 10 ms cloud processing plus the selected communication delay. Their similarity should not be generalized into a network performance claim because only one trace and one fixed phase are used. They merely show that an informative positive is admitted early when tokens are unconstrained.')
p('At the 300 ms deadline, the 50 requests/s development condition changes from a miss to a timely confirmation. This threshold crossing illustrates why papers should report continuous delay as well as pass or fail. A binary table alone would conceal that the confirmation occurs at 223.6 ms and is therefore much closer to 300 ms than to 150 ms.')
p('At 1 request/s, development and test both miss 300 ms by wide margins. The eventual confirmations demonstrate that the forest can identify an admitted attack event, but the admission process exposes it too late for the tested deadlines. This is a scheduling result under the chosen event sequence, not evidence that the forest requires several seconds to classify an individual frame.')
p('The local-warning delay remains 5 ms because an episode-associated positive occurs immediately under the chosen threshold. That favorable timing coexists with an event-level positive fraction near 5% on normal traffic. In operational terms, the local model trades a prompt episode-associated warning for a stream of benign warnings. Reporting either result alone would give an incomplete picture.')

subheading('False Alert Exposure Analysis')
p('The grouped local rates of 3568.1 and 3620.3 alerts/h are close to the theoretical behavior of the one-second anchored grouping rule under dense positives. They should not be interpreted as independent operator decisions occurring at exactly that rate; they quantify how persistently the warning state is re-entered under the selected definition. Alternative grouping, suppression, or cooldown rules would change the number and must be specified before comparing systems.')
p('Development contains 44 288 benign local positives, while test contains 5068. These raw counts are driven by both event frequency and exposure. Dividing by message counts yields 5.07% and 4.74%, close to the development operating target. The similarity suggests threshold transfer at the frame level for these two recordings, but the much shorter test exposure prevents a strong stability conclusion.')
p('The forest produces no observed level-2 alert on either normal recording across all budgets. Because lower-budget requests are subsets selected by time and token availability, the budget rows are not independent trials. Nor does testing five budgets multiply the normal exposure. The appropriate statement is zero observed confirmations over 438.2 s of normal traffic under the evaluated replay configurations.')
p('A one-sided statistical upper bound could be computed only after defining an event process and independent exposure units. The grouped alerts are temporally dependent and the recordings are not random samples from a specified population. The manuscript therefore avoids a confidence interval for the zero count and instead reports the duration directly. This makes the evidence inspectable without implying a population rate that the design cannot support.')
p('The two alert levels imply different operational burdens. If level 1 is visible to a driver or operator, the forest does not remove the high warning frequency. If level 1 is logged silently and only level 2 is surfaced, the earliest actionable signal becomes the confirmation and inherits its budget-dependent delay. A deployment study must define which output triggers which action before false-alert reduction can be translated into benefit.')

subheading('Communication Exposure Analysis')
p('Finite-budget request counts track nominal capacity over the combined recording duration, with small deviations caused by the initial bucket and candidate availability. Development sends 422, 2095, 4186, and 20 889 requests at budgets of 1, 5, 10, and 50 requests/s. Test sends 67, 326, 650, and 3071. These totals provide an audit of admission but do not measure bytes or airtime.')
p('Unrestricted admission sends 49 539 development requests and 6503 test requests, corresponding to approximately 118.3 and 100.5 requests/s. The difference reflects the frequency of local positives in both the normal and attack files. It also shows that the local gate does not impose a small communication footprint at the selected threshold; more than one hundred candidate events per second can remain.')
p('At 50 requests/s, the request stream is approximately half of the unrestricted rate, yet test retains the same 31.0 ms confirmation delay. Development, in contrast, changes from 29.0 to 223.6 ms. This contrast illustrates the main systems result: communication savings and timing loss depend on where informative events fall relative to token availability. A single global budget cannot be justified from these two trajectories.')
p('A byte-level evaluation would require serialization format, feature payload, protocol overhead, authentication material, acknowledgement behavior, and retransmissions. A server evaluation would require batch size, concurrency, queue discipline, and model service time under load. The request counts are therefore an admission metric, not a network-capacity or cloud-cost estimate.')
p('The normal local-positive stream can consume most available tokens before an attack event arrives. The present simulator continuously refills the bucket and does not reserve capacity for state changes, priority identifiers, or burst detection. A future admission policy could incorporate uncertainty, temporal novelty, message identifier, or episode context, but it would need frozen selection rules and comparison against simple baselines.')

subheading('Claims Supported by the Current Evidence')
p('The results support four direct statements. Both evaluated episodes receive an immediate level-1 warning under the assumed 5 ms local time. The local level is extremely noisy on the two normal recordings. Level-2 confirmations produce no observed false alerts on those recordings. Confirmation timeliness varies sharply with budget and recording, including failures at 150 ms under finite development budgets.')
p('The results do not support a population detection rate, a recommended production budget, a claim that cloud execution outperforms local execution, or a safety guarantee. They also do not establish that the forest is more accurate than every lightweight alternative in the literature. These boundaries are part of the result, because they identify the measurements needed to progress from a replay case study to an engineering evaluation.')

subheading('Scenario Based Interpretation')
p('Consider first a monitoring-only deployment in which level-1 warnings are recorded locally and level-2 confirmations are sent to a security operations center. In that setting, the high local warning rate may be tolerable if storage and processing are sufficient, while the absence of observed cloud false alerts could reduce analyst-facing noise. The relevant deadline may be seconds rather than hundreds of milliseconds. The present replay can inform request provisioning, but it does not measure analyst workload or storage cost.')
p('Consider instead a driver-facing advisory. A level-1 warning every second would be unusable, so either the local threshold, grouping logic, or presentation policy must change. Surfacing only confirmations could reduce visible false alerts, but development confirmation would miss 150 ms under every finite budget. The architecture would need an independently justified advisory deadline and human-factors validation before either output could be exposed.')
p('For automated containment, the evidence is weaker. Acting on level 1 risks frequent unnecessary interventions; waiting for level 2 introduces communication dependence and can miss a strict deadline. A graded response might apply a reversible local precaution followed by confirmed escalation, but the safety and availability consequences must be analyzed for the controlled subsystem. The replay contains no actuation model and cannot select such a policy.')
p('For forensic assistance, eventual confirmation may be valuable even when late. The 1 request/s condition eventually identifies both evaluated episodes, demonstrating that severe rate limitation does not necessarily erase all evidence. A forensic system could prioritize completeness over immediate response and batch contextual features. That use case would require different metrics, including evidence retention, attribution quality, and time to analyst review.')
p('For fleet learning, individual confirmations could become labels for later model updates. This creates feedback risks: cloud errors may reinforce the local model, selective querying can bias the observed sample, and connectivity differs across vehicles. Federated architectures such as [16] address collaborative training but not this per-alert selection bias. The current experiment does not update either model online.')
p('These scenarios show why one budget cannot be declared optimal without an operational objective. A budget that is inadequate for rapid containment may be sufficient for forensic confirmation. A threshold that is useful for silent telemetry may be unacceptable for human alerts. The manuscript therefore reports the response surface and evidence boundaries rather than selecting a universal configuration.')

subheading('Threshold and Calibration Implications')
p('The 95th-percentile local threshold is easy to reproduce and produces a deliberately permissive gate. Its event-level behavior transfers approximately from development normal traffic to the short test normal recording, but the grouped burden remains extreme. A practical calibration would specify a target such as alerts per hour, missed-episode tolerance, or request consumption, then tune on multiple development sessions that represent expected operation.')
p('Probability calibration matters if the local score will drive uncertainty routing. Logistic regression scores may be calibrated under its training distribution, but preprocessing, class balance, regularization, and dataset shift can change their meaning. Calibration should be assessed with reliability diagrams and proper scoring rules on independent recordings. Percentile thresholds alone define ranking positions, not comparable probabilities across vehicles.')
p('The forest output is currently used as a hard class. A confirmation threshold could trade false confirmations against episode coverage and should be selected independently of the local gate. Joint calibration could define three outcomes: confirm, reject, or defer for more context. Such a design would more closely resemble reliability routing [14], but it would also require more data than the two episodes available here.')
p('Threshold comparison should use common constraints. One approach is to match local models at a fixed grouped false-alert rate and compare timely episode coverage. Another is to match request rate and compare confirmation delay. A third is to construct a Pareto frontier across workload, communication, and missed episodes. These analyses avoid favoring a model simply because it operates at a more permissive point.')
p('Temporal smoothing can reduce warnings without changing the underlying classifier. Consecutive positives, majority windows, per-identifier cooldowns, or change-point logic can suppress isolated events. They also add detection delay and may merge distinct events. Any smoothing rule should be included in the timing path and frozen before final evaluation rather than applied only to improve a figure.')

subheading('Network and Service Engineering Implications')
p('The request budget is implemented at the event source, but real systems may enforce limits at a gateway, modem, broker, or cloud endpoint. Placement affects which traffic competes for capacity and how bursts are handled. A vehicle-wide gateway may need to arbitrate among multiple buses and applications. The current per-recording bucket represents only the intrusion stream and therefore understates shared-resource complexity.')
p('End-to-end service time should be decomposed into serialization, uplink scheduling, propagation, gateway traversal, queueing, inference, response serialization, downlink, and local handling. CICV5G supplies a recorded communication component, while the replay adds a constant cloud-processing term. Calibration is needed to ensure that processing or application delay is not counted twice and that the trace endpoint matches the proposed architecture.')
p('A finite cloud service can become unstable when arrival rate approaches throughput. Unrestricted local-positive rates above 100 requests/s may be manageable for one vehicle but scale quickly across a fleet. Batching can increase throughput while adding waiting time; prioritization can protect urgent events while starving routine requests. A deployment evaluation should report queue utilization and tail latency, not only average service time.')
p('Loss and disconnection require explicit policies. Requests may be retried, stored, abandoned, or handled locally. Retries increase load and can arrive after relevance has expired. A deadline-aware requester could suppress attempts whose predicted completion is too late, but its prediction must account for queue state and uncertainty. The current token bucket has no feedback from network condition.')
p('Privacy and minimization also shape payload design. Sending decoded signals, identifiers, or raw frames may expose vehicle behavior. Feature compression can reduce bandwidth and disclosure but may prevent the cloud model from reproducing local context. Encryption and authentication add bytes and processing. These concerns are outside the numerical results yet central to judging the feasibility of cloud confirmation.')


heading('Discussion and Limitations')
subheading('Interpretation and Operational Implications')
p('The results support a narrow interpretation: an early local indication and a later confirmed indication expose different latency and false-alert trade-offs. Cloud confirmation cannot be assumed to improve the earliest warning time, since the warning already exists. A false negative at the local stage also prevents a cloud query for that message, which limits the second stage’s ability to recover locally missed activity. Confirmation should therefore not be equated with an independent parallel detector.')

p('A practical interface could distinguish an unconfirmed advisory from a confirmed escalation, but this experiment does not measure operator behavior. If all local warnings trigger the same intervention, the false-alert burden remains dominant; if warnings are hidden until confirmed, the system inherits confirmation delays. The alert policy is part of the system specification.')

subheading('Internal and External Validity')
p('The evaluation includes only two evaluated attack episodes from the same attack family and vehicle context, with one episode per split. Repeating deadline calculations or budget settings does not increase the number of independent attack observations. Broader conclusions require additional attack families, vehicles, normal driving exposure, and genuinely untouched recordings. The previously inspected test split and the development-tuned local threshold further constrain the strength of generalization claims.')
p('The network trace is replayed independently of the CAN recordings, with a fixed initial phase and wraparound. There is no measured coupling between vehicle state, payload size, network load, and intrusion traffic. No remote server or vehicle gateway is exercised. The model also omits processing contention, so unconstrained concurrency can be optimistic. Varying trace phase, communication conditions, and processing assumptions is necessary before interpreting deadline compliance as robust.')
subheading('Computation and Baseline Requirements')
p('A saved development-machine benchmark reports single-message inference of approximately 0.22 ms for logistic regression and 16.0 ms for the forest. The latter exceeds the replay’s assumed 10 ms cloud processing time. These observations motivate sensitivity analysis, but cannot establish automotive hardware requirements. A forest executed locally is an important baseline before arguing that the cloud is preferable; classifier replacement and threshold calibration should also be compared under a common false-alert constraint.')

p('Within this queue-free model, adding 6 ms to cloud service would shift every confirmation by 6 ms without changing admission order. This is an algebraic sensitivity illustration, not an additional measured run; queues or different hardware could change the outcome substantially.')
subheading('Next Validation Steps')
p('The next evaluation should freeze the two-level configuration separately from the earlier policy study, record model and data hashes, and retain per-request selection and response logs. Multiple trace phases and independent normal sessions should test timing and false-alert stability. Local-only operation, a forest executed locally, and alternative admission policies should be compared at matched alert or communication constraints. These extensions would distinguish a useful confirmation mechanism from improvements achievable through local calibration alone.')

subheading('Alternative Architectures')
p('The two-level design is only one point in a broader architecture space. A strong-local configuration would execute the forest on the vehicle and remove communication delay while increasing local compute. An always-cloud configuration would query every frame and test the value of the local gate. A parallel configuration would run both models independently, allowing the cloud path to recover local false negatives. A cascade with abstention would request help for uncertain local scores rather than only positive predictions.')
p('Each alternative changes the meaning of a result. Strong-local performance tests whether offloading is necessary. Always-cloud performance establishes an admission upper bound but may be infeasible. Parallel detection changes coverage because cloud evaluation no longer depends on local positives. Uncertainty routing can reduce requests, but only if scores are calibrated and uncertainty correlates with errors. These baselines should be evaluated under the same data, deadlines, and false-alert definitions.')
p('The present local-positive gate is attractive because it is simple and auditable. It also aligns communication with the events most likely to produce a warning. Its weakness is equally clear: a missed local event cannot be recovered, and a noisy threshold consumes capacity. The architecture is therefore best viewed as a transparent baseline for selective confirmation rather than a final routing policy.')
p('A queue could preserve rejected candidates, but it would require expiration and priority rules. An episode-aware aggregator could compress multiple related positives into one request. A change detector could reserve capacity for sudden shifts. A learned router could optimize a cost function involving delay, false alerts, and requests. Each addition creates parameters and training requirements that must be evaluated without using the test recordings for selection.')

subheading('Validity Threats')
p('Construct validity depends on the mapping between saved labels and the operational concept of a masquerade attack. The ROAD attack traces are post-processed research artifacts. Episode onset follows labels and the 0.5 s gap rule, not an independent physical ground truth. A different episode definition could change first-detection delay and the number of evaluated events.')
p('Metric validity depends on alert semantics. Grouped false alerts use a one-second window anchored at each group start. This is reproducible but not universal. A user-interface cooldown, stateful alarm, or per-identifier suppression rule could produce a different workload. The manuscript reports raw positives and grouped alerts together so readers can reconstruct the relationship.')
p('Internal validity is constrained by the fixed trace phase and assumed processing constants. The budget effects are deterministic consequences of event order, token availability, forest output, and replayed delay. Without controlled perturbations, the experiment cannot attribute a particular delay change to one component. The reported comparisons are descriptive rather than causal estimates.')
p('External validity is constrained by one vehicle context, one masquerade family, two evaluated episodes, two normal recordings, and one cellular trace. Temperature, road condition, driver behavior, firmware, signal-decoding quality, network cell, and service load may alter both detector and communication behavior. Generalization must be tested through independent sessions rather than inferred from frame count.')
p('Conclusion validity is constrained by the lack of independent attack units. Averaging 0/1 outcomes across deadlines or budgets would create pseudo-replication. The paper therefore reports each split and deadline condition explicitly. Statistical comparison among routing policies is deferred until multiple attacks, vehicles, and trace phases are available.')
p('Reproducibility validity depends on software and artifact provenance. Model serialization can be sensitive to library versions, and the original environment file did not pin all relevant packages. The model file has been renamed to reflect its LogisticRegression class, and conflicting result directories have been separated. A release should still provide exact versions and hashes before another team attempts reproduction.')

subheading('Safety and Human Factors')
p('An intrusion alert is not itself a mitigation. A vehicle system must decide whether to log, notify, isolate a bus segment, enter a degraded mode, or request human intervention. Each action has consequences if the alert is false, late, or missing. The present experiment deliberately stops before that decision boundary and therefore cannot claim improved vehicle safety.')
p('The severe level-1 warning rate would be unacceptable for many human-facing interfaces. Repeated alarms can produce habituation, distraction, or disabled monitoring. Cloud confirmation may support escalation, but only if unconfirmed local warnings are handled differently. Human-factors testing is required to determine whether the two levels communicate useful uncertainty or create confusion.')
p('Automation changes the risk balance. Acting on every local warning would inherit the false-alert burden, while waiting for confirmation could delay response beyond an application deadline. A fail-operational design may combine local containment with later confirmation, but the containment action must be justified through hazard analysis. No such control policy is evaluated here.')
p('Communication failure must also be represented explicitly. A missing confirmation can mean no token, packet loss, service outage, late reply, or negative forest output. These cases have different operational meanings but collapse into the absence of a level-2 alert in the current interface. A deployment protocol should expose status and reason codes rather than treating silence as a benign decision.')
p('Security of the assistance channel is outside the simulator. Authentication, confidentiality, integrity, replay protection, denial of service, and model-service compromise could affect both availability and trust. A cloud-confirmation architecture expands the attack surface and must not assume that remote output is inherently authoritative.')

subheading('Reproducibility and Release Plan')
p('The repository now identifies one canonical command for model training and prediction generation and one for the two-level replay. The replay stores its effective configuration beside the result tables, and separate output directories prevent unified-policy and cloud-confirmation experiments from silently overwriting the paper evidence. A release should additionally record file hashes and the exact software environment.')
p('The release manifest should contain hashes for raw source files where redistribution is permitted, processed prediction files, models, scaler, delay trace, scripts, tables, and figures. If raw datasets cannot be redistributed, the manifest should give the official source, expected filename, checksum, and preprocessing command. This supports provenance without copying restricted data.')
p('Per-request logs should record split, recording, event identifier, observation time, local score, local prediction, token state, admission decision, selected delay sample, forest output, reply time, and deadline classification. These fields would make every plotted point traceable to events and help distinguish selection delay from communication delay.')
p('Software dependencies should be pinned to exact versions in a lock file or container specification. Serialized scikit-learn models should be loaded with the same supported version used for training, or exported to a stable interchange format after equivalence testing. Hardware, operating system, processor, and thread settings should accompany any performance benchmark.')
p('The manuscript, Word source, final PDF, figures, and numerical tables should be generated from the same frozen results. IEEE Access requires matching Word and PDF submissions. A release check should compare table cells and plotted values against the saved summary, scan citations, and render every page before submission.')

subheading('Prioritized Experimental Roadmap')
p('The first priority is more independent data. Additional masquerade recordings should span attack targets, durations, intensities, vehicles, and operating conditions. Longer normal sessions are needed to characterize rare confirmations and workload. A new untouched evaluation split should be reserved before threshold or policy development resumes.')
p('The second priority is baseline control. Logistic regression, the forest, and at least one lightweight sequence or signal method should be compared locally at matched false-alert constraints. Cloud assistance should then be evaluated against strong-local, always-cloud, random, periodic, uncertainty-based, and local-positive admission under the same deadlines and communication accounting.')
p('The third priority is systems realism. Replay should sample multiple trace phases and network contexts, model packet loss and retransmission, enforce a finite service queue, and measure target-hardware processing. Payload bytes, cryptographic overhead, energy, and server concurrency should be reported alongside requests. These additions would test whether a policy remains timely when capacity is shared.')
p('The fourth priority is prespecified analysis. Thresholds, deadlines, episode rules, grouping windows, trace phases, and primary metrics should be frozen before the final holdout is inspected. Confidence intervals or hypothesis tests should use independent episodes or recordings as units. Exploratory results can remain useful, but they should be labeled and separated from confirmatory claims.')
p('The fifth priority is operational validation. Engineers and domain experts should define the meaning of warning and confirmation, the actions attached to each level, and the maximum acceptable delay. Human-factors evaluation should measure comprehension and alarm burden. Safety analysis should address false interventions, missed attacks, communication loss, and malicious manipulation of the cloud path.')

subheading('Lifecycle and Governance Considerations')
p('A deployed detector will experience software updates, vehicle aging, repairs, sensor replacement, seasonal conditions, and changes in driver behavior. These factors can shift signal distributions and request volume. Monitoring should therefore track local score distributions, warning rates, confirmation rates, request utilization, and tail latency by vehicle and software version. A stable accuracy estimate at release time is not sufficient for lifecycle assurance.')
p('Threshold changes are operational model changes even when classifier weights remain fixed. Lowering the local threshold can increase coverage and cloud demand; raising it can reduce burden while blocking confirmation opportunities. Threshold revisions should use version control, documented development evidence, staged rollout, and rollback criteria. Production data used for tuning must be separated from the next independent evaluation set.')
p('Cloud-model updates create another compatibility boundary. The local gate, feature encoder, and cloud forest must agree on identifiers, signal ordering, missing-value representation, and scaling. A server update that changes feature expectations can silently turn accepted requests into invalid inputs. Interface schemas, model versions, and validation checks should accompany each request and response.')
p('Logging is necessary for audit but can expose sensitive vehicle behavior. Retention should be limited to fields required for security analysis, with access control, integrity protection, and deletion rules. Event identifiers should support traceability without unnecessarily identifying a driver. The current research files do not define a production data-governance policy.')
p('Fleet deployment also requires fairness across heterogeneous vehicles and connectivity conditions. Vehicles with poor coverage may receive fewer timely confirmations, while models trained on one platform may generate more warnings on another. Reporting only pooled fleet metrics could hide these disparities. Evaluation should stratify by vehicle platform, environment, software version, and network condition where sample sizes permit.')
p('Incident response must account for disagreement between levels. A local positive with cloud negative, a cloud timeout, and a cloud positive are distinct states. Policies should define who receives each state, how long it remains active, and what evidence is retained. The present architecture records local warnings and positive confirmations, but a production protocol should expose negative and unavailable outcomes explicitly.')
p('Governance should also specify when cloud assistance may be disabled. Maintenance mode, privacy restrictions, cost limits, network outage, or suspected service compromise may require local-only operation. The system should degrade predictably and indicate that confirmation is unavailable. Because the local detector is noisy in this study, local-only fallback would require a different threshold or response policy before operational use.')
p('These lifecycle requirements reinforce the paper’s central distinction: a second prediction is only one component of an alerting system. Evidence about classifiers, scheduling, communication, interfaces, and governance must remain linked. Expanding the experiment along only one dimension, such as adding a larger cloud model, would not by itself establish that the complete system is effective or safe.')

p('Finally, no physical actuation is performed. The system produces warning and confirmation logs, and the deadlines are evaluation parameters. A safety claim would require an application-specific response requirement, a defined action policy, target-hardware measurements, and validation beyond this replay study.')

heading('Conclusion')
p('V2C-Sentinel separates immediate local warnings from budget-limited cloud confirmation in a chronological CAN masquerade replay. Both evaluated episodes receive a local warning after the assumed 5 ms processing time. Cloud confirmation shows no false alerts on the two short normal recordings, but fails to meet 150 ms in several finite-budget conditions. The evidence therefore favors reporting the two alert levels separately and treating confirmation timeliness as conditional on the recording and communication budget. Additional independent data, stronger local baselines, and end-to-end timing measurements are needed before making deployment or safety claims.')
p('Acknowledgment', 'Heading 5')
p('OpenAI Codex was used to generate and revise the abstract and Sections I–VII, organize the reported repository results, prepare the reference list and figure-generation code, and format this manuscript. Figures 1–4 were produced with this assistance; numerical plots use saved experimental outputs. No new experimental measurements were generated for the manuscript. The authors remain responsible for verification of the final text, figures, and references.')
p('References', 'Heading 5')
for ref in ['M. E. Verma et al., “A comprehensive guide to CAN IDS data and introduction of the ROAD dataset,” PLOS ONE, vol. 19, no. 1, Art. no. e0296879, 2024, doi: 10.1371/journal.pone.0296879.', 'M. H. Shahriar, Y. Xiao, P. Moriano, W. Lou, and Y. T. Hou, “CANShield: Deep-learning-based intrusion detection framework for controller area networks at the signal level,” IEEE Internet Things J., vol. 10, no. 24, pp. 22111–22127, 2023, doi: 10.1109/JIOT.2023.3303271.', 'X. Zhang et al., “5G communication delay dataset for cloud-based vehicle planning and control,” Scientific Data, vol. 13, Art. no. 878, 2026, doi: 10.1038/s41597-026-07239-7.', 'P. Moriano, R. A. Bridges, and M. D. Iannacone, “Detecting CAN masquerade attacks with signal clustering similarity,” in Proc. Workshop Automotive and Autonomous Vehicle Security (AutoSec), 2022. [Online]. Available: https://www.ndss-symposium.org/wp-content/uploads/autosec2022_23028_paper.pdf', 'W. Marfo, P. Moriano, D. K. Tosh, and S. V. Moore, “Detecting masquerade attacks in controller area networks using graph machine learning,” IEEE Trans. Inf. Forensics Security, vol. 20, pp. 13127–13142, 2025, doi: 10.1109/TIFS.2025.3636019.', 'P. Moriano, S. C. Hespeler, M. Li, and R. A. Bridges, “Evaluating lightweight unsupervised online IDS for masquerade attacks in CAN,” J. Inf. Secur. Appl., vol. 98, Art. no. 104392, 2026, doi: 10.1016/j.jisa.2026.104392.', 'Q. Liu, R. Song, L. Cui, H. Zhang, Y. Sun, and L. Sun, “MIDS: Detecting stealthy masquerade and tampering attacks on CAN bus via bidirectional Mamba,” arXiv:2606.18599, 2026, doi: 10.48550/arXiv.2606.18599.', 'L. Guerra et al., “AI-driven intrusion detection systems (IDS) on the ROAD dataset: A comparative analysis for automotive controller area network (CAN),” in Proc. Cyber Security in CarS Workshop (CSCS), 2024, pp. 39–49, doi: 10.1145/3689936.3694696.', 'S. Rajapaksha, H. Kalutarage, M. O. Al-Kadri, A. Petrovski, and G. Madzudzo, “Beyond vanilla: Improved autoencoder-based ensemble in-vehicle intrusion detection system,” J. Inf. Secur. Appl., vol. 77, Art. no. 103570, 2023, doi: 10.1016/j.jisa.2023.103570.', 'T.-N. Hoang and D. Kim, “Detecting in-vehicle intrusion via semi-supervised learning-based convolutional adversarial autoencoders,” Veh. Commun., vol. 38, Art. no. 100520, 2022, doi: 10.1016/j.vehcom.2022.100520.', 'C. Dong, H. Wu, and Q. Li, “Multiple observation HMM-based CAN bus intrusion detection system for in-vehicle network,” IEEE Access, vol. 11, pp. 35639–35648, 2023, doi: 10.1109/ACCESS.2023.3265018.', 'G. Loukas, T. Vuong, R. Heartfield, G. Sakellari, Y. Yoon, and D. Gan, “Cloud-based cyber-physical intrusion detection for vehicles using deep learning,” IEEE Access, vol. 6, pp. 3491–3508, 2018, doi: 10.1109/ACCESS.2017.2782159.', 'S. Chinchali et al., “Network offloading policies for cloud robotics: A learning-based approach,” arXiv:1902.05703, 2019, doi: 10.48550/arXiv.1902.05703.', 'S. Azam, F. Naaz, and M. M. Salim, “A reliability-aware edge–cloud framework for early intrusion detection in IoT networks,” Electronics, vol. 15, no. 16, Art. no. 3506, 2026, doi: 10.3390/electronics15163506.', 'A. Mao, M. Mohri, and Y. Zhong, “Mastering multiple-expert routing: Realizable H-consistency and strong guarantees for learning to defer,” in Proc. 42nd Int. Conf. Machine Learning, PMLR, vol. 267, 2025, pp. 43035–43066.', 'A. Heidari, S. H. Rastegar, and A. Khonsari, “Explainable edge–cloud federated intrusion detection system for electric vehicles using GRUs, DRL selection, and integrated gradients,” J. Big Data, early access, Aug. 2026, doi: 10.1186/s40537-026-01505-6.']: p(ref, 'references')

biography(
    'Asmaa Berdigh',
    'received the engineering degree in embedded systems and industrial computing from the National School of Applied Sciences of Fez, Morocco, and the Ph.D. degree from Moulay Ismail University, Meknes, Morocco, in 2024.',
    'She has professional experience in automotive research and development. Her research interests include the Internet of Things, connected vehicles, intra-vehicle networks, wireless communication, and automotive engineering. Her publications address connected-car architectures, vehicle-to-everything communication, and hybrid wired-to-wireless in-vehicle communication.',
)

biography(
    'Khalid El Yassini',
    'received the degree in applied mathematics, with a specialization in statistics, from Abdelmalek Essaadi University, Morocco, in 1991, and the M.Sc. and Ph.D. degrees in mathematics, with a specialization in operations research, from the University of Sherbrooke, Canada, in 1994 and 2000, respectively.',
    'He is a Professor of applied mathematics and computer science with the Faculty of Sciences, Moulay Ismail University, Meknes, Morocco. He has taught mathematics, computer science, operations research, and logistics at institutions in Canada and Morocco. His research interests include mathematical programming, artificial intelligence, information systems, computer-network security, telecommunications, intelligent systems, and logistics engineering.',
)

final_sect=deepcopy(sections[1])
body.append(final_sect)
# Preserve source styling; add only manuscript body, image relationships and media.
from io import BytesIO
buffer=BytesIO(); d.save(buffer); buffer.seek(0)
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','word/footer1.xml'}
with ZipFile(REF) as original, ZipFile(buffer) as revised, ZipFile(OUT,'w',ZIP_DEFLATED) as out:
    for item in original.infolist():
        if item.filename == 'word/footer1.xml':
            data=original.read(item.filename).replace(b'2017', b'2026')
            data=data.replace(
                b'<w:t>1</w:t>',
                b'<w:fldChar w:fldCharType="begin"/><w:instrText xml:space="preserve"> PAGE </w:instrText><w:fldChar w:fldCharType="separate"/><w:t>1</w:t><w:fldChar w:fldCharType="end"/>',
                1,
            )
        else:
            data=revised.read(item.filename) if item.filename in editable else original.read(item.filename)
        out.writestr(item,data)
    for name in revised.namelist():
        if name not in original.namelist():
            assert name.startswith('word/media/'),name
            out.writestr(name,revised.read(name))
with ZipFile(REF) as original, ZipFile(OUT) as out:
    changes=[n for n in original.namelist() if original.read(n)!=out.read(n)]
assert set(changes)<=editable, changes
(ROOT/'paper/V2C_Sentinel_IEEE_17_Page_Journal_Draft.md').write_text('\n\n'.join(source),encoding='utf-8')
print(json.dumps({'output':str(OUT),'words':len(' '.join(source).split()),'changed_parts':changes}))
