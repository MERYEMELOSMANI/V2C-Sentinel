from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'paper/working'
REF = WORK / 'ieee-conference-a4.docx'
OUT = ROOT / 'paper/V2C_Sentinel_IEEE_Draft.docx'
d = Document(REF)
sections = [deepcopy(s._sectPr) for s in d.sections]
body = d._element.body
for child in list(body): body.remove(child)
source = []

def p(text, style='Body Text'):
    para = d.add_paragraph(text, style)
    source.append(text)
    return para

def heading(text): return p(text, 'Heading 1')

def boundary(sect):
    para = d.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = Pt(1)
    para._p.get_or_add_pPr().append(sect)

title = 'V2C Sentinel for Local Warnings and Cloud Confirmation of CAN Masquerade Attacks'
p(title, 'paper title')
p('[Author names]\n[Department and institution]\n[City, country]\n[Email or ORCID]', 'Author').paragraph_format.space_after = Pt(10)
title_sect = deepcopy(sections[0])
for node in list(title_sect):
    if node.tag in (qn('w:footerReference'), qn('w:headerReference'), qn('w:titlePg')): title_sect.remove(node)
boundary(title_sect)

p('Abstract—Cloud assistance can provide a second intrusion-detection decision, but communication delay and request limits affect when that decision becomes available. This paper examines a two-level alert architecture for controller area network masquerade detection. A local logistic regression model issues an immediate warning and, subject to a token-bucket budget, requests confirmation from a random forest model. A chronological replay combines automotive attack recordings with a separately recorded fifth-generation cellular delay trace. The evaluation contains one attack episode and one normal recording in each of the development and test splits. Local warnings detect both episodes after the assumed 5 ms processing time, while producing approximately 3568 and 3620 grouped false alerts per hour on the respective normal recordings. Cloud confirmation produces no observed false alerts on those normal recordings, but its timeliness depends on the request budget and recording. At 50 requests per second, confirmation takes 223.6 ms on development and 31.0 ms on test; unrestricted requests reduce these delays to 29.0 and 31.0 ms. These observations support separating early warnings from confirmed alerts, while demonstrating that confirmation cannot be assumed to meet a 150 ms deadline. The study is exploratory and does not establish deployment performance or population-level detection rates.', 'Abstract')
p('Keywords—controller area network, intrusion detection, cloud assistance, detection latency, request budgeting', 'Keywords')

heading('Introduction')
p('Controller area network (CAN) intrusion detection must distinguish malicious activity from legitimate vehicle traffic. The ROAD dataset includes masquerade attack traces designed to support this investigation [1]. Signal-level methods such as CANShield use relationships among signals to detect attacks that may be difficult to identify from message timing alone [2]. These studies motivate examining both the detection decision and the time at which an actionable warning becomes available.')
p('Consulting a second detector introduces a separate systems question. A local model can produce an early indication, whereas remote analysis requires a request, processing, and a returning response. Cellular measurements in the CICV5G dataset provide recorded vehicle–network–vehicle delays for studying this communication component [3]. A constrained request budget can introduce additional waiting before an informative message is selected. Consequently, classifier output alone does not determine episode-level alert latency.')
p('This paper asks how separating immediate local warnings from cloud-confirmed alerts affects false-alert burden and episode-level timeliness under limited request rates. V2C-Sentinel preserves each local positive decision as a level-1 warning and sends eligible messages for a level-2 decision. Confirmation is an additional output; a negative or delayed remote decision does not retract the earlier warning.')
p('The contribution is a reproducible replay case study of this separation, with a shared token-bucket mechanism and distinct measurements for the two alert levels. We report development and test recordings separately and expose the conditions under which confirmation misses an experimental deadline. We do not propose a new classifier or claim that remote execution is necessary. Both detectors execute on the development computer, and the communication channel is simulated from recorded delays.')

heading('Related Work')
p('Verma et al. [1] introduce ROAD and discuss the characteristics of CAN intrusion-detection data. The selected correlated-signal masquerade traces in this study come from that dataset. Their use supplies labeled recordings for replay, but does not make separate recordings equivalent to independent vehicles or attack families. In particular, the masquerade traces must not be described as evidence of an end-to-end physical attack experiment conducted by this study.')
p('Shahriar et al. [2] present CANShield, a signal-level intrusion-detection framework using multiple autoencoders and an ensemble decision. Its focus is learning temporal and signal relationships. Our experiment fixes simpler supervised classifiers and examines the delivery of local and confirmed alerts. No direct comparison with CANShield is made, so the results do not establish superior classifier accuracy or detection responsiveness.')
p('Zhang et al. [3] publish CICV5G measurements for cloud-based vehicle planning and control. We reuse one urban n78 trace as an exogenous communication-delay sequence. This application does not provide synchronized CAN and network measurements, nor does it validate an intrusion-detection service over a live cellular link. The present comparison is deliberately limited to two alert levels within one replay architecture; a comprehensive comparison with vehicle offloading and selective-inference methods remains future work.')

heading('System and Replay Method')
p('Each input event contains a timestamp, message identifier, extracted signal values, and an evaluation label. The label is used only for training or evaluation as appropriate to the split. The prediction pipeline uses the identifier and Signal_ columns as features, replaces missing values with −1, and standardizes features using statistics fitted on the training split. Timestamp and label columns are excluded from the feature matrix.')
p('The local classifier is logistic regression with a maximum of 1000 iterations and random seed 42. The remote classifier is a random forest with 50 trees, maximum depth 15, and random seed 42. Both classifiers are trained on the normal and attack training recordings. The local positive threshold is the 95th percentile of attack-class probabilities on the normal development recording; equality to the threshold is classified as positive. The forest uses its predicted class. This local threshold is distinct from the suspicion threshold used in an earlier policy-selection experiment.')
p('For a message observed at time t, local processing completes at t + 5 ms. A positive local prediction immediately produces a level-1 warning. Only locally positive messages are eligible for a cloud request. A token bucket replenishes at b requests/s, has capacity max(2, 0.1b) requests, and starts full for each recording. One admitted request consumes one token. Budgets of 1, 5, 10, and 50 requests/s are compared with an unrestricted condition. Unrestricted means that every locally positive message is requested, rather than every CAN message.')
p('An admitted request returns after the local completion time plus the replayed communication delay and 10 ms of assumed cloud processing. Positive forest predictions produce level-2 alerts at their return times. Replies are ordered chronologically, and replies outstanding at the end of a recording are processed. A negative forest prediction produces no level-2 alert. The simulator does not model a service queue, packet loss, retransmissions, or concurrent inference capacity.')
p('Trace timestamps are measured relative to the first trace sample. At each request time, the simulator wraps time by the trace duration and selects the first trace sample at or after the wrapped time. Replayed delays have a 1 ms minimum. The recorded delay is used directly as the communication-delay term; the two datasets have no common time alignment. Whether the measurement includes application processing that overlaps the added 10 ms term requires calibration before an end-to-end deployment interpretation.')
p('The reporting deadlines are 150 and 300 ms from attack onset. They classify observed detection times and do not suppress replies or change scheduling. In particular, the local warning is not held until a deadline, and late confirmations remain available as diagnostic information. These deadlines are experimental settings, not validated automotive safety requirements.')

heading('Experimental Setup and Metrics')
p('The ROAD training pair consists of the basic-long ambient recording and correlated-signal masquerade recording 1. Development uses the radio-infotainment ambient recording and masquerade recording 2. Test uses the winter ambient recording and masquerade recording 3. These correspond to normal_01 and attack_01, normal_02 and attack_02, and normal_03 and attack_03 in the saved split manifest. The cellular replay uses the CICV5G urban-road n78 trace. Table I reports the evaluated recordings; the durations are the last timestamp minus the first timestamp.')

def table(caption, headers, rows, widths):
    cap = p(caption, 'table head')
    # Explicit caption text avoids inheriting sample automatic numbering.
    n = OxmlElement('w:numPr'); v=OxmlElement('w:numId'); v.set(qn('w:val'),'0'); n.append(v); cap._p.get_or_add_pPr().append(n)
    cap.paragraph_format.keep_with_next = True
    tbl = d.add_table(rows=1, cols=len(headers)); tbl.autofit=False
    tbl.alignment=1
    for col,w in zip(tbl.columns,widths): col.width=Inches(w)
    for cell,text in zip(tbl.rows[0].cells,headers): cell.text=text
    for row in rows:
        for cell,text in zip(tbl.add_row().cells,row): cell.text=str(text)
    pr=tbl._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'000000'); borders.append(edge)
    pr.append(borders)
    for i,row in enumerate(tbl.rows):
        trp=row._tr.get_or_add_trPr(); trp.append(OxmlElement('w:cantSplit'))
        if i==0: trp.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j])
            for para in cell.paragraphs:
                para.style=d.styles['table copy']; para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before=Pt(2); para.paragraph_format.space_after=Pt(2)
                for run in para.runs: run.bold=(i==0)
    source.append(caption+'\n'+' | '.join(headers)+'\n'+'\n'.join(' | '.join(map(str,r)) for r in rows))

table('TABLE I.  EVALUATION RECORDINGS', ['Recording','Messages','Duration (s)','Episodes'], [
    ['Development attack', '63 260','28.227','1'], ['Development normal','874 018','390.456','0'],
    ['Test attack','38 003','16.964','1'],['Test normal','106 939','47.731','0']], [1.17,.75,.76,.65])
p('An attack episode groups positive-labeled events whose successive timestamps are separated by no more than 0.5 s. Each evaluated attack recording contains one such episode. Detection delay is the earliest warning time linked to a true-positive event in that episode minus episode onset. A detection is timely when this delay is no greater than the reporting deadline. Because each split contains one episode, timely results are reported as detected episode counts rather than as population estimates.')
p('False-alert rates in the main comparison use only the separate normal recordings. Warning times are sorted, the first false warning starts an alert group, and a new group begins when a subsequent warning is more than 1 s after the start of the current group. Thus, the grouping window is anchored at each group start, rather than extended by every warning. The grouped count divided by recording duration in hours gives grouped false alerts per hour. Under dense warnings this rule approaches one grouped alert per second; short-recording boundary effects can yield a rate slightly above 3600 per hour.')
p('Request totals are summed over the attack and normal recordings within each split, once per budget. They are not summed across deadlines because the same events are evaluated at both deadlines. Raw totals should not be compared across splits without accounting for their different durations. The token bucket permits an initial burst, so a nominal budget specifies a sustained admission rate rather than a strict cap in every one-second interval.')
p('The test recordings were inspected during earlier development. The recorded protocol states that they were not used for model fitting or threshold selection, but they are not an untouched holdout. The present two-level experiment is a later analysis than the earlier frozen multi-policy protocol. Accordingly, these results should be treated as exploratory. Processing times of 5 and 10 ms are simulation assumptions, not measurements on vehicle hardware.')

heading('Results')
p('Level 1 detects the single episode in each split after 5.0 ms under every budget, satisfying both reporting deadlines. This invariance follows from the architecture: the local output does not wait for cloud admission or response. The normal development recording produces 3568.1 grouped local false alerts per hour; the normal test recording produces 3620.3. These values indicate a persistently noisy first level, despite successful early warnings on the two attack episodes.')
p('Table II reports first cloud-confirmation delays. At the 150 ms deadline, none of the four finite budgets confirms the development episode on time. The 50 requests/s condition becomes timely when the deadline increases to 300 ms. On test, 5, 10, and 50 requests/s meet both deadlines, while 1 request/s misses both. Unrestricted requests meet both deadlines in both splits. Every listed condition eventually confirms its evaluated episode.')
table('TABLE II.  CLOUD CONFIRMATION DELAY', ['Budget\n(requests/s)','Development\n(ms)','Test\n(ms)'],[
    ['1','3207.9','725.8'],['5','805.7','129.8'],['10','703.6','129.8'],['50','223.6','31.0'],['Unrestricted','29.0','31.0']], [1.13,1.1,1.1])
p('The absence of a test-delay improvement between 5 and 10 requests/s shows that additional requests do not necessarily select an earlier informative event in a particular replay. Conversely, increasing the development budget to 50 requests/s improves confirmation delay substantially but still does not meet 150 ms. These are recording-specific observations; the experiment does not isolate the causal contributions of token availability, prediction sequence, and trace phase.')
p('No level-2 false alerts are observed on either normal recording at any budget. The total normal exposure is approximately 438.2 s, or 7.3 min. Zero observed alerts over this short exposure does not establish a zero underlying false-alert rate. In addition, level-1 warnings remain present, so the combined architecture does not eliminate the local warning burden. An operational reduction in operator workload would depend on how the two outputs are presented and acted upon.')
table('TABLE III.  CLOUD REQUEST TOTALS BY SPLIT', ['Budget\n(requests/s)','Development','Test'],[
    ['1','422','67'],['5','2095','326'],['10','4186','650'],['50','20889','3071'],['Unrestricted','49539','6503']], [1.13,1.1,1.1])
p('Table III combines requests from the attack and normal recording in each split. These counts measure communication demand in requests, not bytes, monetary cost, or link utilization. The normal traffic consumes tokens as well as producing false local warnings. Consequently, sustained false positives can compete with attack-related messages for confirmation capacity.')

heading('Discussion and Limitations')
p('The results support a narrow interpretation: an early local indication and a later confirmed indication expose different latency and false-alert trade-offs. Cloud confirmation cannot be assumed to improve the earliest warning time, since the warning already exists. A false negative at the local stage also prevents a cloud query for that message, which limits the second stage’s ability to recover locally missed activity. Confirmation should therefore not be equated with an independent parallel detector.')
p('The evaluation includes only two evaluated attack episodes from the same attack family and vehicle context, with one episode per split. Repeating deadline calculations or budget settings does not increase the number of independent attack observations. Broader conclusions require additional attack families, vehicles, normal driving exposure, and genuinely untouched recordings. The previously inspected test split and the development-tuned local threshold further constrain the strength of generalization claims.')
p('The network trace is replayed independently of the CAN recordings, with a fixed initial phase and wraparound. There is no measured coupling between vehicle state, payload size, network load, and intrusion traffic. No remote server or vehicle gateway is exercised. The model also omits processing contention, so unconstrained concurrency can be optimistic. Varying trace phase, communication conditions, and processing assumptions is necessary before interpreting deadline compliance as robust.')
p('A saved development-machine benchmark reports single-message inference of approximately 0.22 ms for logistic regression and 16.0 ms for the forest. The latter exceeds the replay’s assumed 10 ms cloud processing time. These observations motivate sensitivity analysis, but cannot establish automotive hardware requirements. A forest executed locally is an important baseline before arguing that the cloud is preferable; classifier replacement and threshold calibration should also be compared under a common false-alert constraint.')
p('Finally, no physical actuation is performed. The system produces warning and confirmation logs, and the deadlines are evaluation parameters. A safety claim would require an application-specific response requirement, a defined action policy, target-hardware measurements, and validation beyond this replay study.')

heading('Conclusion')
p('V2C-Sentinel separates immediate local warnings from budget-limited cloud confirmation in a chronological CAN masquerade replay. Both evaluated episodes receive a local warning after the assumed 5 ms processing time. Cloud confirmation shows no false alerts on the two short normal recordings, but fails to meet 150 ms in several finite-budget conditions. The evidence therefore favors reporting the two alert levels separately and treating confirmation timeliness as conditional on the recording and communication budget. Additional independent data, stronger local baselines, and end-to-end timing measurements are needed before making deployment or safety claims.')
p('Acknowledgment', 'Heading 5')
p('OpenAI Codex was used to generate and revise the abstract and Sections I–VII, organize the reported repository results, prepare the reference list, and format this manuscript. This assistance did not generate new experimental measurements. [Authors to verify the manuscript, finalize this disclosure, and add any applicable funding acknowledgment.]')
p('References', 'Heading 5')
for ref in [
    'M. E. Verma et al., “A comprehensive guide to CAN IDS data and introduction of the ROAD dataset,” PLOS ONE, vol. 19, no. 1, Art. no. e0296879, 2024, doi: 10.1371/journal.pone.0296879.',
    'M. H. Shahriar, Y. Xiao, P. Moriano, W. Lou, and Y. T. Hou, “CANShield: Deep learning-based intrusion detection framework for controller area networks at the signal-level,” arXiv:2205.01306, ver. 4, Oct. 2023, doi: 10.48550/arXiv.2205.01306.',
    'X. Zhang et al., “5G communication delay dataset for cloud-based vehicle planning and control,” Scientific Data, vol. 13, Art. no. 878, 2026, doi: 10.1038/s41597-026-07239-7.'
]: p(ref, 'references')

final_sect=deepcopy(sections[3])
body.append(final_sect)
# Preserve all template package parts except the manuscript body and placeholder footer.
from lxml import etree
document_xml=etree.tostring(d._element,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(REF) as original, ZipFile(OUT,'w',ZIP_DEFLATED) as out:
    for item in original.infolist():
        data=original.read(item.filename)
        if item.filename=='word/document.xml': data=document_xml
        out.writestr(item,data)
with ZipFile(REF) as original, ZipFile(OUT) as out:
    changes=[n for n in original.namelist() if original.read(n)!=out.read(n)]
assert changes==['word/document.xml'], changes
(ROOT/'paper/V2C_Sentinel_IEEE_Draft.md').write_text('\n\n'.join(source),encoding='utf-8')
print(json.dumps({'output':str(OUT),'words':len(' '.join(source).split()),'changed_parts':changes}))
