from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'paper/working'
REF = WORK / 'ieee-conference-a4.docx'
OUT = ROOT / 'paper/V2C_Sentinel_IEEE_Full_Paper.docx'
d = Document(REF)
sections = [deepcopy(s._sectPr) for s in d.sections]
body = d._element.body
for child in list(body): body.remove(child)
source = []

def p(text, style='Body Text'):
    para = d.add_paragraph(text, style)
    source.append(text)
    return para

def heading(text): return p(text, 'Heading 1')

def boundary(sect):
    para = d.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = Pt(1)
    para._p.get_or_add_pPr().append(sect)


def subheading(text): return p(text, 'Heading 2')

def figure(name,caption):
    para=d.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(5)
    para.paragraph_format.space_after=Pt(2)
    para.paragraph_format.keep_with_next=True
    shape=para.add_run().add_picture(str(ROOT/'paper/figures'/name),width=Inches(3.34))
    shape._inline.docPr.set('descr',caption)
    cap=p(caption,'figure caption')
    cap.paragraph_format.keep_together=True
    cap.paragraph_format.space_after=Pt(7)
    num=OxmlElement('w:numPr'); ident=OxmlElement('w:numId'); ident.set(qn('w:val'),'0'); num.append(ident)
    cap._p.get_or_add_pPr().append(num)

def equation(tokens,number):
    para=d.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before=Pt(4)
    para.paragraph_format.space_after=Pt(6)
    math=OxmlElement('m:oMath')
    def run(text):
        r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=text; r.append(t); return r
    for token in tokens:
        if isinstance(token,tuple):
            sub=OxmlElement('m:sSub'); e=OxmlElement('m:e'); e.append(run(token[0])); v=OxmlElement('m:sub'); v.append(run(token[1])); sub.append(e); sub.append(v); math.append(sub)
        else: math.append(run(token))
    para._p.append(math); para.add_run('    ('+str(number)+')')
    source.append('Equation '+str(number)+': '+str(tokens))

title = 'V2C Sentinel for Local Warnings and Cloud Confirmation of CAN Masquerade Attacks'
p(title, 'paper title')
p('[Author names]\n[Department and institution]\n[City, country]\n[Email or ORCID]', 'Author').paragraph_format.space_after = Pt(10)
title_sect = deepcopy(sections[0])
for node in list(title_sect):
    if node.tag in (qn('w:footerReference'), qn('w:headerReference'), qn('w:titlePg')): title_sect.remove(node)
boundary(title_sect)

p('Abstract—Cloud assistance can provide a second intrusion-detection decision, but communication delay and request limits affect when that decision becomes available. This paper examines a two-level alert architecture for controller area network masquerade detection. A local logistic regression model issues an immediate warning and, subject to a token-bucket budget, requests confirmation from a random forest model. A chronological replay combines automotive attack recordings with a separately recorded fifth-generation cellular delay trace. The evaluation contains one attack episode and one normal recording in each of the development and test splits. Local warnings detect both episodes after the assumed 5 ms processing time, while producing approximately 3568 and 3620 grouped false alerts per hour on the respective normal recordings. Cloud confirmation produces no observed false alerts on those normal recordings, but its timeliness depends on the request budget and recording. At 50 requests per second, confirmation takes 223.6 ms on development and 31.0 ms on test; unrestricted requests reduce these delays to 29.0 and 31.0 ms. These observations support separating early warnings from confirmed alerts, while demonstrating that confirmation cannot be assumed to meet a 150 ms deadline. The study is exploratory and does not establish deployment performance or population-level detection rates.', 'Abstract')
p('Keywords—controller area network, intrusion detection, cloud assistance, detection latency, request budgeting', 'Keywords')

heading('Introduction')
p('Controller area network (CAN) intrusion detection must distinguish malicious activity from legitimate vehicle traffic. The ROAD dataset includes masquerade attack traces designed to support this investigation [1]. Signal-level methods such as CANShield use relationships among signals to detect attacks that may be difficult to identify from message timing alone [2]. These studies motivate examining both the detection decision and the time at which an actionable warning becomes available.')
p('Consulting a second detector introduces a separate systems question. A local model can produce an early indication, whereas remote analysis requires a request, processing, and a returning response. Cellular measurements in the CICV5G dataset provide recorded vehicle–network–vehicle delays for studying this communication component [3]. A constrained request budget can introduce additional waiting before an informative message is selected. Consequently, classifier output alone does not determine episode-level alert latency.')
p('This paper asks how separating immediate local warnings from cloud-confirmed alerts affects false-alert burden and episode-level timeliness under limited request rates. V2C-Sentinel preserves each local positive decision as a level-1 warning and sends eligible messages for a level-2 decision. Confirmation is an additional output; a negative or delayed remote decision does not retract the earlier warning.')
p('The contribution is a reproducible replay case study of this separation, with a shared token-bucket mechanism and distinct measurements for the two alert levels. We report development and test recordings separately and expose the conditions under which confirmation misses an experimental deadline. We do not propose a new classifier or claim that remote execution is necessary. Both detectors execute on the development computer, and the communication channel is simulated from recorded delays.')

p('The evaluation distinguishes local-warning timeliness, false-alert burden at each alert level, and the confirmation budget needed to meet an episode-onset deadline. This prevents a low confirmation false-alert count from being mistaken for a low total warning burden.')


p('Related Work', 'Heading 1')
p('CAN Masquerade Detection and Timing', 'Heading 2')
p('ROAD [1] provides the data foundation for this study. Its masquerade traces are post-processed from recorded vehicle data, and should not be represented as organically executed masquerade attacks. CANShield [2] uses signal-level autoencoder ensembles and explicitly evaluates attack and event detection latency as well as hardware processing latency. Thus, measuring time from an attack to a detection is not itself a new contribution. Our focus is the additional timing and admission cost of obtaining a second model decision.', 'Body Text')
p('Moriano et al. [4] study signal-clustering similarity, Marfo et al. [5] use temporally enriched message-sequence graphs, and Moriano et al. [6] compare lightweight online methods. Together, they motivate separating classifier evidence, observation-window requirements, and response delivery rather than reporting one undifferentiated latency.', 'Body Text')
p('The MIDS preprint by Liu et al. [7] uses bidirectional Mamba processing for CAN identifiers and payloads and reports single-window inference latency. Guerra et al. [8] compare multiple intrusion detectors on ROAD and other automotive data. These results provide relevant detector and dataset context, but their performance figures cannot be transferred to the present classifiers, split, or hardware. We therefore do not place their reported accuracy beside our episode counts as if the experiments were directly comparable.', 'Body Text')
p('Rajapaksha et al. [9], Hoang and Kim [10], and Dong et al. [11] offer different local CAN detectors. Their results show that local model design and efficiency remain active alternatives to cloud assistance; remote inference should not be assumed necessary.', 'Body Text')
p('Vehicle Cloud Offloading and Selective Assistance', 'Heading 2')
p('Loukas et al. [12] investigate cloud-based cyber-physical intrusion detection using deep models on a small robotic vehicle. Their study includes a latency model for deciding when computation offloading is beneficial. This is a close architectural predecessor: cloud-based vehicle intrusion detection and latency-aware offloading are established ideas. V2C-Sentinel differs in evaluating selected CAN masquerade recordings, preserving an immediate local warning, and measuring the first confirmed alert under a finite request-admission budget.', 'Body Text')
p('Chinchali et al. [13] formulate cloud-robotics offloading as a sequential decision problem balancing performance against communication cost. The present fixed local-positive admission rule and token bucket do not learn when cloud assistance is beneficial; this is a constrained replay evaluation, not a general routing method.', 'Body Text')
p('Reliability Routing and Federated Detection', 'Heading 2')
p('Azam et al. [14] present a reliability-aware edge–cloud intrusion-detection framework that uses calibrated packet-prefix evidence to choose local decisions, continued observation, or cloud refinement. This is a close conceptual comparison because cloud access is selective. V2C-Sentinel does not implement calibrated abstention or a continue-observing action. It forwards locally positive CAN events when tokens are available and then evaluates confirmation against elapsed-time deadlines. The difference is in routing semantics and evaluation setting, not in the existence of an edge–cloud structure.', 'Body Text')
p('Mao et al. [15] formalize cost-aware routing among experts. Heidari et al. [16] select participants for collaborative vehicle-model training. Neither mechanism is equivalent to this fixed, per-message cloud-confirmation replay.', 'Body Text')
p('Scope of the Present Contribution', 'Heading 2')
p('CICV5G [3] supplies an independently recorded cellular delay trace. Combining that trace with ROAD allows a controlled study of how a request budget changes confirmation time while the local output remains available. The evidence supports a small, deterministic case study rather than a first-of-its-kind cloud detector. In particular, the deadline is used to evaluate alerts; the implemented request policy does not predict whether a reply will meet that deadline. Comparative accuracy, routing optimality, and deployment safety remain outside the demonstrated claims.', 'Body Text')

heading('System and Replay Method')

subheading('Threat Model and Alert Semantics')
p('The replay assumes an adversary capable of generating masquerading CAN traffic represented by the selected ROAD labels. The attacker can alter the content associated with legitimate traffic; the study does not simulate initial compromise, cryptographic authentication, or a live attacker. The detector observes recorded identifiers and extracted signals. Model files, local processing, and the conceptual cloud service are trusted. Attacks on the confirmation channel, adversarial model manipulation, and intentional denial of the request budget are outside the evaluated threat model.')
p('A level-1 warning means that the local score crossed its operating threshold. A level-2 confirmation means that an admitted message received a positive forest prediction. The latter is a model output, not proof of an attack. Absence of confirmation can mean that the message was not admitted, that its reply has not arrived, or that the forest predicted benign traffic. Fig. 1 shows these separate paths and the persistence of the local warning.')
figure('fig1_architecture.png','Fig. 1. Two-level replay architecture. Positive local predictions produce warnings immediately and compete for cloud admission. Recorded communication delays affect only the confirmation path.')
subheading('Feature Processing and Fixed Classifiers')

p('Each input event contains a timestamp, message identifier, extracted signal values, and an evaluation label. The label is used only for training or evaluation as appropriate to the split. The prediction pipeline uses the identifier and Signal_ columns as features, replaces missing values with −1, and standardizes features using statistics fitted on the training split. Timestamp and label columns are excluded from the feature matrix.')
p('The local classifier is logistic regression with a maximum of 1000 iterations and random seed 42. The remote classifier is a random forest with 50 trees, maximum depth 15, and random seed 42. Both classifiers are trained on the normal and attack training recordings. The local positive threshold is the 95th percentile of attack-class probabilities on the normal development recording; equality to the threshold is classified as positive. The forest uses its predicted class. This local threshold is distinct from the suspicion threshold used in an earlier policy-selection experiment.')
subheading('Admission Control and Response Timing')
p('For a message observed at time t, local processing completes at t + 5 ms. A positive local prediction immediately produces a level-1 warning. Only locally positive messages are eligible for a cloud request. A token bucket replenishes at b requests/s, has capacity max(2, 0.1b) requests, and starts full for each recording. One admitted request consumes one token. Budgets of 1, 5, 10, and 50 requests/s are compared with an unrestricted condition. Unrestricted means that every locally positive message is requested, rather than every CAN message.')
p('An admitted request returns after the local completion time plus the replayed communication delay and 10 ms of assumed cloud processing. Positive forest predictions produce level-2 alerts at their return times. Replies are ordered chronologically, and replies outstanding at the end of a recording are processed. A negative forest prediction produces no level-2 alert. The simulator does not model a service queue, packet loss, retransmissions, or concurrent inference capacity.')

p('Let tᵢ denote the observation time of message i, τL the assumed local processing time, τC the assumed cloud processing time, and d(s) the replay delay at request time s. For a locally positive message, the warning time is given by (1). If that message is admitted, its confirmation time, conditional on a positive forest prediction, is given by (2).')
equation([('s','i'),' = ',('t','i'),' + ',('τ','L')],1)
equation([('c','i'),' = ',('s','i'),' + d(',('s','i'),') + ',('τ','C')],2)
p('Equations (1) and (2) show why an episode-level deadline is stricter than a per-request response target: a suitable message may be admitted only after the attack has already been active. There is no waiting queue for rejected requests. An eligible message without a token is skipped, and a later local positive may obtain the next token. Thus the observed confirmation delay combines event selection after onset with the delay of the selected reply.')

subheading('Delay Replay and Deadline Interpretation')
p('Trace timestamps are measured relative to the first trace sample. At each request time, the simulator wraps time by the trace duration and selects the first trace sample at or after the wrapped time. Replayed delays have a 1 ms minimum. The recorded delay is used directly as the communication-delay term; the two datasets have no common time alignment. Whether the measurement includes application processing that overlaps the added 10 ms term requires calibration before an end-to-end deployment interpretation.')
p('The reporting deadlines are 150 and 300 ms from attack onset. They classify observed detection times and do not suppress replies or change scheduling. In particular, the local warning is not held until a deadline, and late confirmations remain available as diagnostic information. These deadlines are experimental settings, not validated automotive safety requirements.')

heading('Experimental Setup and Metrics')
subheading('Data Partition and Experimental Conditions')
p('The ROAD training pair consists of the basic-long ambient recording and correlated-signal masquerade recording 1. Development uses the radio-infotainment ambient recording and masquerade recording 2. Test uses the winter ambient recording and masquerade recording 3. These correspond to normal_01 and attack_01, normal_02 and attack_02, and normal_03 and attack_03 in the saved split manifest. The cellular replay uses the CICV5G urban-road n78 trace. Table I reports the evaluated recordings; the durations are the last timestamp minus the first timestamp.')

def table(caption, headers, rows, widths):
    cap = p(caption, 'table head')
    # Explicit caption text avoids inheriting sample automatic numbering.
    n = OxmlElement('w:numPr'); v=OxmlElement('w:numId'); v.set(qn('w:val'),'0'); n.append(v); cap._p.get_or_add_pPr().append(n)
    cap.paragraph_format.keep_with_next = True
    tbl = d.add_table(rows=1, cols=len(headers)); tbl.autofit=False
    tbl.alignment=1
    for col,w in zip(tbl.columns,widths): col.width=Inches(w)
    for cell,text in zip(tbl.rows[0].cells,headers): cell.text=text
    for row in rows:
        for cell,text in zip(tbl.add_row().cells,row): cell.text=str(text)
    pr=tbl._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'000000'); borders.append(edge)
    pr.append(borders)
    for i,row in enumerate(tbl.rows):
        trp=row._tr.get_or_add_trPr(); trp.append(OxmlElement('w:cantSplit'))
        if i==0: trp.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j])
            for para in cell.paragraphs:
                para.style=d.styles['table copy']; para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before=Pt(2); para.paragraph_format.space_after=Pt(2)
                for run in para.runs: run.bold=(i==0)
    source.append(caption+'\n'+' | '.join(headers)+'\n'+'\n'.join(' | '.join(map(str,r)) for r in rows))

table('TABLE I.  EVALUATION RECORDINGS', ['Recording','Messages','Duration (s)','Episodes'], [
    ['Development attack', '63 260','28.227','1'], ['Development normal','874 018','390.456','0'],
    ['Test attack','38 003','16.964','1'],['Test normal','106 939','47.731','0']], [1.17,.75,.76,.65])
subheading('Episode and Alert Metrics')
p('An attack episode groups positive-labeled events whose successive timestamps are separated by no more than 0.5 s. Each evaluated attack recording contains one such episode. Detection delay is the earliest warning time linked to a true-positive event in that episode minus episode onset. A detection is timely when this delay is no greater than the reporting deadline. Because each split contains one episode, timely results are reported as detected episode counts rather than as population estimates.')

p('For episode e with onset aₑ and alert level k, let Wₑ,ₖ be the set of alert times linked to positive-labeled events in that episode. Its first-detection delay is defined by (3), with the minimum of an empty set treated as infinity. Timeliness at deadline D is the condition Δₑ,ₖ ≤ D. This definition credits an alert associated with the episode even if the reply arrives after the episode has ended; its delay is still measured from onset.')
equation([('Δ','e,k'),' = min{w − ',('a','e'),': w ∈ ',('W','e,k'),'}'],3)

p('False-alert rates use separate normal recordings. The first false warning starts a group; a new group begins when a later warning is more than 1 s after that group start. Grouped count divided by recording duration gives grouped false alerts per hour. Dense warnings can approach one group per second, and short-recording boundaries can yield a rate slightly above 3600 per hour.')
p('Request totals are summed over the attack and normal recordings within each split, once per budget. They are not summed across deadlines because the same events are evaluated at both deadlines. Raw totals should not be compared across splits without accounting for their different durations. The token bucket permits an initial burst, so a nominal budget specifies a sustained admission rate rather than a strict cap in every one-second interval.')

p('For exposure normalization, Nq denotes admitted requests and T denotes the combined duration in seconds of the normal and attack recording within a split. The pooled request rate is Rq = Nq/T. This is a ratio of totals, not an unweighted average of per-recording rates. Development has 418.683 s of combined exposure and test has 64.694 s. Reporting pooled rates makes communication demand comparable despite the large difference in raw recording length.')
subheading('Reproducibility and Evidence Boundaries')
p('The paper uses only `results/final_two_level/`, whose 40 recording-level evaluations cover two splits, two recordings per split, five budgets, and two reporting deadlines. A deadline change reclassifies the same alert trajectory and is not an independent experiment. The deterministic replay therefore reports observed values without confidence intervals.')

p('The test recordings were inspected during earlier development. The recorded protocol states that they were not used for model fitting or threshold selection, but they are not an untouched holdout. The present two-level experiment is a later analysis than the earlier frozen multi-policy protocol. Accordingly, these results should be treated as exploratory. Processing times of 5 and 10 ms are simulation assumptions, not measurements on vehicle hardware.')

heading('Results')
subheading('Early Warnings and Confirmation Timeliness')
p('Level 1 produces its first episode-associated warning after 5.0 ms in each split, satisfying both reporting deadlines. This does not wait for cloud admission or response. The normal development and test recordings produce 3568.1 and 3620.3 grouped local false alerts per hour, respectively, showing that the first level remains persistently noisy.')
p('Fig. 2 compares first-detection delays, and Table II reports the corresponding deadline outcomes. At the 150 ms deadline, none of the four finite budgets confirms the development episode on time. The 50 requests/s condition becomes timely when the deadline increases to 300 ms. On test, 5, 10, and 50 requests/s meet both deadlines, while 1 request/s misses both. Unrestricted requests meet both deadlines in both splits. Every listed condition eventually confirms its evaluated episode. In budget order 1, 5, 10, 50, and unrestricted, development delays are 3207.9, 805.7, 703.6, 223.6, and 29.0 ms; test delays are 725.8, 129.8, 129.8, 31.0, and 31.0 ms.')
figure('fig2_latency.png','Fig. 2. First-detection delay versus request budget. Each confirmation series contains one attack episode per split. The dotted line is the assumed 5 ms local path; horizontal lines mark 150 and 300 ms. The vertical axis is logarithmic.')
table('TABLE II.  TIMELY CLOUD DETECTIONS OUT OF ONE EPISODE', ['Budget\n(req/s)','Dev\n150 ms','Dev\n300 ms','Test\n150 ms','Test\n300 ms'],[
    ['1','0/1','0/1','0/1','0/1'],['5','0/1','0/1','1/1','1/1'],['10','0/1','0/1','1/1','1/1'],['50','0/1','1/1','1/1','1/1'],['Unrestricted','1/1','1/1','1/1','1/1']], [1.05,.57,.57,.57,.57])
p('The absence of a test-delay improvement between 5 and 10 requests/s shows that additional requests do not necessarily select an earlier informative event in a particular replay. Conversely, increasing the development budget to 50 requests/s improves confirmation delay substantially but still does not meet 150 ms. These are recording-specific observations; the experiment does not isolate the causal contributions of token availability, prediction sequence, and trace phase.')
subheading('False Alerts and the Meaning of Confirmation')
p('No level-2 false alerts are observed on either normal recording at any budget. The total normal exposure is approximately 438.2 s, or 7.3 min. Zero observed alerts over this short exposure does not establish a zero underlying false-alert rate. In addition, level-1 warnings remain present, so the combined architecture does not eliminate the local warning burden. An operational reduction in operator workload would depend on how the two outputs are presented and acted upon.')

figure('fig3_false_alerts.png','Fig. 3. False-alert burden on the separate normal recordings. Rates are identical across tested budgets. Zero confirmed alerts were observed over only 438.19 s combined exposure; local warnings remain present.')
p('The normal development file contains 44 288 positive local predictions among 874 018 events, while the normal test file contains 5068 among 106 939 events. These correspond to event-level positive fractions of approximately 5.07% and 4.74%. The grouping rule transforms these dense warning streams into 387 and 48 grouped alerts, respectively. These counts explain why the grouped rates remain near one alert per second. Event-level false positives and grouped operator alerts are therefore different quantities and should not be used interchangeably.')
subheading('Communication Demand')

table('TABLE III.  CLOUD REQUEST TOTALS BY SPLIT', ['Budget\n(requests/s)','Development','Test'],[
    ['1','422','67'],['5','2095','326'],['10','4186','650'],['50','20889','3071'],['Unrestricted','49539','6503']], [1.13,1.1,1.1])
p('Table III combines requests from the attack and normal recording in each split. These counts measure communication demand in requests, not bytes, monetary cost, or link utilization. The normal traffic consumes tokens as well as producing false local warnings. Consequently, sustained false positives can compete with attack-related messages for confirmation capacity.')

figure('fig4_request_rates.png','Fig. 4. Actual request rates pooled over the normal and attack recording within each split. Initial token bursts and finite exposure allow small deviations from the nominal sustained budget. Unrestricted admission still queries only local positives.')
p('In the unrestricted condition, development sends 49 539 requests over 418.683 s, approximately 118.3 requests/s, and test sends 6503 over 64.694 s, approximately 100.5 requests/s. At a 50 requests/s budget, the totals correspond to approximately 49.9 and 47.5 requests/s. The smaller test rate is consistent with candidate availability as well as admission limits; the budget is not a requirement to transmit at that rate. A deployed design would also need a byte budget and a server-capacity limit.')
p('The requests demonstrate a trade-off rather than a uniformly favorable budget. Development requires unrestricted admission to meet 150 ms, whereas test meets that deadline with 5 requests/s. Selecting a recommended budget from the favorable test recording would overstate the evidence. A useful configuration must instead be selected against independently measured operating requirements and a broader distribution of attack onsets and benign demand.')


heading('Discussion and Limitations')
subheading('Interpretation and Operational Implications')
p('The results support a narrow interpretation: an early local indication and a later confirmed indication expose different latency and false-alert trade-offs. Cloud confirmation cannot be assumed to improve the earliest warning time, since the warning already exists. A false negative at the local stage also prevents a cloud query for that message, which limits the second stage’s ability to recover locally missed activity. Confirmation should therefore not be equated with an independent parallel detector.')

p('A practical interface could distinguish an unconfirmed advisory from a confirmed escalation, but this experiment does not measure operator behavior. If all local warnings trigger the same intervention, the false-alert burden remains dominant; if warnings are hidden until confirmed, the system inherits confirmation delays. The alert policy is part of the system specification.')

subheading('Internal and External Validity')
p('The evaluation includes only two evaluated attack episodes from the same attack family and vehicle context, with one episode per split. Repeating deadline calculations or budget settings does not increase the number of independent attack observations. Broader conclusions require additional attack families, vehicles, normal driving exposure, and genuinely untouched recordings. The previously inspected test split and the development-tuned local threshold further constrain the strength of generalization claims.')
p('The network trace is replayed independently of the CAN recordings, with a fixed initial phase and wraparound. There is no measured coupling between vehicle state, payload size, network load, and intrusion traffic. No remote server or vehicle gateway is exercised. The model also omits processing contention, so unconstrained concurrency can be optimistic. Varying trace phase, communication conditions, and processing assumptions is necessary before interpreting deadline compliance as robust.')
subheading('Computation and Baseline Requirements')
p('A saved development-machine benchmark reports single-message inference of approximately 0.22 ms for logistic regression and 16.0 ms for the forest. The latter exceeds the replay’s assumed 10 ms cloud processing time. These observations motivate sensitivity analysis, but cannot establish automotive hardware requirements. A forest executed locally is an important baseline before arguing that the cloud is preferable; classifier replacement and threshold calibration should also be compared under a common false-alert constraint.')

p('Within this queue-free model, adding 6 ms to cloud service would shift every confirmation by 6 ms without changing admission order. This is an algebraic sensitivity illustration, not an additional measured run; queues or different hardware could change the outcome substantially.')
subheading('Next Validation Steps')
p('The next evaluation should freeze the two-level configuration separately from the earlier policy study, record model and data hashes, and retain per-request selection and response logs. Multiple trace phases and independent normal sessions should test timing and false-alert stability. Local-only operation, a forest executed locally, and alternative admission policies should be compared at matched alert or communication constraints. These extensions would distinguish a useful confirmation mechanism from improvements achievable through local calibration alone.')

p('Finally, no physical actuation is performed. The system produces warning and confirmation logs, and the deadlines are evaluation parameters. A safety claim would require an application-specific response requirement, a defined action policy, target-hardware measurements, and validation beyond this replay study.')

heading('Conclusion')
p('V2C-Sentinel separates immediate local warnings from budget-limited cloud confirmation in a chronological CAN masquerade replay. Both evaluated episodes receive a local warning after the assumed 5 ms processing time. Cloud confirmation shows no false alerts on the two short normal recordings, but fails to meet 150 ms in several finite-budget conditions. The evidence therefore favors reporting the two alert levels separately and treating confirmation timeliness as conditional on the recording and communication budget. Additional independent data, stronger local baselines, and end-to-end timing measurements are needed before making deployment or safety claims.')
p('Acknowledgment', 'Heading 5')
p('OpenAI Codex was used to generate and revise the abstract and Sections I–VII, organize the reported repository results, prepare the reference list and figure-generation code, and format this manuscript. Figures 1–4 were produced with this assistance; numerical plots use saved experimental outputs. No new experimental measurements were generated for the manuscript. The authors remain responsible for verification of the final text, figures, and references.')
p('References', 'Heading 5')
for ref in ['M. E. Verma et al., “A comprehensive guide to CAN IDS data and introduction of the ROAD dataset,” PLOS ONE, vol. 19, no. 1, Art. no. e0296879, 2024, doi: 10.1371/journal.pone.0296879.', 'M. H. Shahriar, Y. Xiao, P. Moriano, W. Lou, and Y. T. Hou, “CANShield: Deep-learning-based intrusion detection framework for controller area networks at the signal level,” IEEE Internet Things J., vol. 10, no. 24, pp. 22111–22127, 2023, doi: 10.1109/JIOT.2023.3303271.', 'X. Zhang et al., “5G communication delay dataset for cloud-based vehicle planning and control,” Scientific Data, vol. 13, Art. no. 878, 2026, doi: 10.1038/s41597-026-07239-7.', 'P. Moriano, R. A. Bridges, and M. D. Iannacone, “Detecting CAN masquerade attacks with signal clustering similarity,” in Proc. Workshop Automotive and Autonomous Vehicle Security (AutoSec), 2022. [Online]. Available: https://www.ndss-symposium.org/wp-content/uploads/autosec2022_23028_paper.pdf', 'W. Marfo, P. Moriano, D. K. Tosh, and S. V. Moore, “Detecting masquerade attacks in controller area networks using graph machine learning,” IEEE Trans. Inf. Forensics Security, vol. 20, pp. 13127–13142, 2025, doi: 10.1109/TIFS.2025.3636019.', 'P. Moriano, S. C. Hespeler, M. Li, and R. A. Bridges, “Evaluating lightweight unsupervised online IDS for masquerade attacks in CAN,” J. Inf. Secur. Appl., vol. 98, Art. no. 104392, 2026, doi: 10.1016/j.jisa.2026.104392.', 'Q. Liu, R. Song, L. Cui, H. Zhang, Y. Sun, and L. Sun, “MIDS: Detecting stealthy masquerade and tampering attacks on CAN bus via bidirectional Mamba,” arXiv:2606.18599, 2026, doi: 10.48550/arXiv.2606.18599.', 'L. Guerra et al., “AI-driven intrusion detection systems (IDS) on the ROAD dataset: A comparative analysis for automotive controller area network (CAN),” in Proc. Cyber Security in CarS Workshop (CSCS), 2024, pp. 39–49, doi: 10.1145/3689936.3694696.', 'S. Rajapaksha, H. Kalutarage, M. O. Al-Kadri, A. Petrovski, and G. Madzudzo, “Beyond vanilla: Improved autoencoder-based ensemble in-vehicle intrusion detection system,” J. Inf. Secur. Appl., vol. 77, Art. no. 103570, 2023, doi: 10.1016/j.jisa.2023.103570.', 'T.-N. Hoang and D. Kim, “Detecting in-vehicle intrusion via semi-supervised learning-based convolutional adversarial autoencoders,” Veh. Commun., vol. 38, Art. no. 100520, 2022, doi: 10.1016/j.vehcom.2022.100520.', 'C. Dong, H. Wu, and Q. Li, “Multiple observation HMM-based CAN bus intrusion detection system for in-vehicle network,” IEEE Access, vol. 11, pp. 35639–35648, 2023, doi: 10.1109/ACCESS.2023.3265018.', 'G. Loukas, T. Vuong, R. Heartfield, G. Sakellari, Y. Yoon, and D. Gan, “Cloud-based cyber-physical intrusion detection for vehicles using deep learning,” IEEE Access, vol. 6, pp. 3491–3508, 2018, doi: 10.1109/ACCESS.2017.2782159.', 'S. Chinchali et al., “Network offloading policies for cloud robotics: A learning-based approach,” arXiv:1902.05703, 2019, doi: 10.48550/arXiv.1902.05703.', 'S. Azam, F. Naaz, and M. M. Salim, “A reliability-aware edge–cloud framework for early intrusion detection in IoT networks,” Electronics, vol. 15, no. 16, Art. no. 3506, 2026, doi: 10.3390/electronics15163506.', 'A. Mao, M. Mohri, and Y. Zhong, “Mastering multiple-expert routing: Realizable H-consistency and strong guarantees for learning to defer,” in Proc. 42nd Int. Conf. Machine Learning, PMLR, vol. 267, 2025, pp. 43035–43066.', 'A. Heidari, S. H. Rastegar, and A. Khonsari, “Explainable edge–cloud federated intrusion detection system for electric vehicles using GRUs, DRL selection, and integrated gradients,” J. Big Data, early access, Aug. 2026, doi: 10.1186/s40537-026-01505-6.']: p(ref, 'references')

final_sect=deepcopy(sections[3])
body.append(final_sect)
# Preserve source styling; add only manuscript body, image relationships and media.
from io import BytesIO
buffer=BytesIO(); d.save(buffer); buffer.seek(0)
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with ZipFile(REF) as original, ZipFile(buffer) as revised, ZipFile(OUT,'w',ZIP_DEFLATED) as out:
    for item in original.infolist():
        data=revised.read(item.filename) if item.filename in editable else original.read(item.filename)
        out.writestr(item,data)
    for name in revised.namelist():
        if name not in original.namelist():
            assert name.startswith('word/media/'),name
            out.writestr(name,revised.read(name))
with ZipFile(REF) as original, ZipFile(OUT) as out:
    changes=[n for n in original.namelist() if original.read(n)!=out.read(n)]
assert set(changes)<=editable, changes
(ROOT/'paper/V2C_Sentinel_IEEE_Full_Paper.md').write_text('\n\n'.join(source),encoding='utf-8')
print(json.dumps({'output':str(OUT),'words':len(' '.join(source).split()),'changed_parts':changes}))
